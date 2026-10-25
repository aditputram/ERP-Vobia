import tempfile
from io import BytesIO
from datetime import date, datetime
from decimal import Decimal
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import User
from audit.models import AuditEvent
from inventory.models import FIFOLayer, InventoryMovement, QCInspection
from master_data.models import Category, Product, ProductStatus, ProductVariant, SKU, Supplier
from merchandising.models import IncomingPlan, ProjectionScenario, SalesProjection
from merchandising.services.workflows import approve_incoming_plan, approve_sales_projection, create_incoming_plan
from production.models import ProductionActivity, ProductionCogsFinalization, ProductionPlan, ProductionStage
from production.services import (
    eligible_activity_choices,
    ensure_production_order,
    submit_cmt_activity_batch,
    submit_production_activity,
)

from .models import (
    PPICProductionMaterial,
    PPICProductionPlan,
    PPICRequirement,
    POWIPImportBatch,
    PurchaseOrder,
)
from .services.wip_import import approve_po_wip_import, create_po_wip_import
from .services.workflows import (
    cancel_po,
    create_draft_po,
    delete_unused_supplier,
    generate_po_from_production_plan,
    release_po,
    revise_legacy_wip_supplier,
    save_ppic_production_plan,
    sync_ppic_requirement,
)


class PurchasingWorkflowTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="buyer", password="test-password")
        status = ProductStatus.objects.create(code="ACTIVE", name="Active")
        category = Category.objects.create(code="APPAREL", name="Apparel")
        product = Product.objects.create(code="P-1", name="Product", status=status, category=category)
        variant = ProductVariant.objects.create(product=product, name="Black", color="Black")
        self.sku = SKU.objects.create(
            sku="SKU-1",
            product_variant=variant,
            current_retail_price=Decimal("200000"),
            current_master_cogs=Decimal("100000"),
        )
        self.supplier = Supplier.objects.create(code="SUP-1", name="Supplier 1")
        self.scenario = ProjectionScenario.objects.create(
            name="September 2026",
            start_month=date(2026, 9, 1),
            end_month=date(2026, 9, 1),
            created_by=self.user,
        )

    def _approved_incoming(self, qty=100):
        projection = SalesProjection.objects.create(
            scenario=self.scenario,
            month=date(2026, 9, 1),
            sku=self.sku,
            system_recommendation=Decimal("100"),
        )
        approve_sales_projection(projection.id, 100, self.user)
        plan = create_incoming_plan(projection.id, 0)
        approve_incoming_plan(plan.id, qty, self.user)
        return IncomingPlan.objects.get(pk=plan.id)

    def test_operation_lists_export_valid_excel_workbooks(self):
        self.client.force_login(self.user)
        self._approved_incoming(100)
        create_draft_po(
            supplier=self.supplier,
            need_month=date(2026, 9, 1),
            actor=self.user,
            manual_lines=[(self.sku, Decimal("5"))],
        )

        requirement_response = self.client.get(reverse("purchasing:requirements"), {"export": "xlsx"})
        self.assertEqual(requirement_response.status_code, 200)
        requirement_workbook = load_workbook(BytesIO(requirement_response.content), read_only=True)
        self.assertEqual(requirement_workbook.sheetnames, ["PPIC Requirement"])
        self.assertEqual(requirement_workbook["PPIC Requirement"]["A1"].value, "Need Month")
        self.assertEqual(requirement_workbook["PPIC Requirement"].max_row, 2)

        po_response = self.client.get(reverse("purchasing:purchase_orders"), {"export": "xlsx"})
        self.assertEqual(po_response.status_code, 200)
        po_workbook = load_workbook(BytesIO(po_response.content), read_only=True)
        self.assertEqual(po_workbook.sheetnames, ["Purchase Order"])
        self.assertEqual(po_workbook["Purchase Order"]["A1"].value, "No. PPO")
        self.assertEqual(po_workbook["Purchase Order"].max_row, 2)

    def test_only_approved_incoming_syncs_and_keeps_revision_history(self):
        plan = self._approved_incoming(100)
        requirement = PPICRequirement.objects.get(incoming_plan=plan)
        self.assertEqual(requirement.approved_qty, Decimal("100"))
        self.assertEqual(requirement.revisions.count(), 1)

        plan.approval_status = IncomingPlan.ApprovalStatus.DRAFT
        plan.save(update_fields=["approval_status"])
        with self.assertRaises(ValidationError):
            sync_ppic_requirement(plan.id, self.user)

    def test_zero_approved_incoming_does_not_create_ppic_requirement(self):
        projection = SalesProjection.objects.create(
            scenario=self.scenario,
            month=date(2026, 9, 1),
            sku=self.sku,
            system_recommendation=Decimal("0"),
        )
        approve_sales_projection(projection.id, 0, self.user)
        plan = create_incoming_plan(projection.id, 0)
        approve_incoming_plan(plan.id, 0, self.user)

        self.assertFalse(PPICRequirement.objects.filter(incoming_plan=plan).exists())

    def test_zero_resync_removes_unallocated_requirement_and_its_revision(self):
        plan = self._approved_incoming(100)
        requirement = PPICRequirement.objects.get(incoming_plan=plan)
        requirement_id = requirement.id
        self.assertEqual(requirement.revisions.count(), 1)

        plan.final_approved_incoming = Decimal("0")
        plan.save(update_fields=["final_approved_incoming"])
        self.assertIsNone(sync_ppic_requirement(plan.id, self.user, "Zero requirement UAT"))

        self.assertFalse(PPICRequirement.objects.filter(pk=requirement_id).exists())

    def test_review_create_release_and_immutable_cogs_snapshot(self):
        plan = self._approved_incoming(100)
        requirement = PPICRequirement.objects.get(incoming_plan=plan)
        po = create_draft_po(
            supplier=self.supplier,
            need_month=date(2026, 9, 1),
            actor=self.user,
            requirement_quantities={requirement.id: 60},
            requirement_cogs={requirement.id: 125000},
        )
        self.assertEqual(po.status, PurchaseOrder.Status.DRAFT)
        self.assertIsNone(po.po_number)
        self.assertEqual(po.lines.get().cogs_snapshot, Decimal("125000"))
        august_release = timezone.make_aware(datetime(2026, 8, 23, 10, 0))
        with patch("purchasing.services.workflows.timezone.now", return_value=august_release):
            po = release_po(po.id, self.user)
        self.assertEqual(po.po_number, "PO-VOB-08/26-001")
        self.assertEqual(po.issue_month, date(2026, 8, 1))
        self.assertEqual(po.need_month, date(2026, 9, 1))
        line = po.lines.get()
        self.assertEqual(line.cogs_snapshot, Decimal("125000"))

        self.sku.current_master_cogs = Decimal("120000")
        self.sku.save(update_fields=["current_master_cogs"])
        line.refresh_from_db()
        self.assertEqual(line.cogs_snapshot, Decimal("125000"))

        second = create_draft_po(
            supplier=self.supplier,
            need_month=date(2026, 9, 1),
            actor=self.user,
            requirement_quantities={requirement.id: 40},
        )
        with patch("purchasing.services.workflows.timezone.now", return_value=august_release):
            self.assertEqual(release_po(second.id, self.user).po_number, "PO-VOB-08/26-002")
        requirement.refresh_from_db()
        self.assertEqual(requirement.remaining_qty, Decimal("0"))

        september_release = timezone.make_aware(datetime(2026, 9, 2, 10, 0))
        third = create_draft_po(
            supplier=self.supplier,
            need_month=date(2026, 9, 1),
            actor=self.user,
            manual_lines=[(self.sku, 1)],
        )
        with patch("purchasing.services.workflows.timezone.now", return_value=september_release):
            third = release_po(third.id, self.user)
        self.assertEqual(third.po_number, "PO-VOB-09/26-003")
        self.assertEqual(third.issue_month, date(2026, 9, 1))

        january_release = timezone.make_aware(datetime(2027, 1, 2, 10, 0))
        fourth = create_draft_po(
            supplier=self.supplier,
            need_month=date(2027, 1, 1),
            actor=self.user,
            manual_lines=[(self.sku, 1)],
        )
        with patch("purchasing.services.workflows.timezone.now", return_value=january_release):
            fourth = release_po(fourth.id, self.user)
        self.assertEqual(fourth.po_number, "PO-VOB-01/27-001")

    def test_cannot_order_above_remaining_requirement(self):
        plan = self._approved_incoming(100)
        requirement = PPICRequirement.objects.get(incoming_plan=plan)
        with self.assertRaises(ValidationError):
            create_draft_po(
                supplier=self.supplier,
                need_month=date(2026, 9, 1),
                actor=self.user,
                requirement_quantities={requirement.id: 101},
            )
        self.assertEqual(PPICRequirement.objects.count(), 1)

    def test_po_generator_hides_requirements_without_saved_production_plan(self):
        plan = self._approved_incoming(100)
        requirement = PPICRequirement.objects.get(incoming_plan=plan)
        self.client.force_login(self.user)
        response = self.client.get(reverse("purchasing:generator"), {"month": "2026-09"})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["plans"], [])
        self.assertContains(response, "Requirement yang belum dibuatkan Production Plan tidak dapat langsung menjadi PO")
        self.assertNotContains(response, requirement.sku.sku)
        self.assertEqual(PurchaseOrder.objects.count(), 0)

    def test_cmt_production_plan_builds_cost_and_generates_draft_po(self):
        incoming = self._approved_incoming(100)
        requirement = PPICRequirement.objects.get(incoming_plan=incoming)
        plan = save_ppic_production_plan(
            actor=self.user,
            requirements=[requirement],
            quantities={str(requirement.id): "100"},
            values={
                "production_method": PPICProductionPlan.ProductionMethod.CMT,
                "supplier": self.supplier,
                "target_material_purchase_date": date(2026, 8, 10),
                "target_cut_start_date": date(2026, 8, 15),
                "target_cut_end_date": date(2026, 8, 16),
                "target_make_start_date": date(2026, 8, 17),
                "target_make_end_date": date(2026, 8, 25),
                "target_trim_start_date": date(2026, 8, 26),
                "target_trim_end_date": date(2026, 8, 28),
                "target_qc_start_date": date(2026, 8, 29),
                "target_qc_end_date": date(2026, 8, 30),
                "target_inbound_date": date(2026, 9, 1),
                "notes": "CMT test",
            },
            materials=[{
                "product": self.sku.product_variant.product,
                "supplier": self.supplier,
                "material_name": "Fabric",
                "quantity": Decimal("200"),
                "unit": "meter",
                "unit_price": Decimal("5000"),
            }],
            process_costs=[{
                "product": self.sku.product_variant.product,
                "cmt_supplier": self.supplier,
                "cmt_unit_cost": Decimal("40000"),
                "additional_process_unit_cost": Decimal("0"),
            }],
        )

        self.assertEqual(plan.material_cost, Decimal("1000000"))
        self.assertEqual(plan.planned_total_cost, Decimal("5000000"))
        self.assertEqual(plan.planned_unit_cogs, Decimal("50000"))
        self.assertEqual(plan.lines.get().unit_cogs, Decimal("50000"))
        self.assertEqual(plan.materials.get().product, self.sku.product_variant.product)
        po = generate_po_from_production_plan(plan.id, self.user)
        plan.refresh_from_db()
        self.assertEqual(plan.status, PPICProductionPlan.Status.GENERATED)
        self.assertEqual(plan.purchase_order_id, po.id)
        self.assertEqual(po.status, PurchaseOrder.Status.DRAFT)
        self.assertEqual(po.child_code, "B")
        self.assertEqual(po.order_type, PurchaseOrder.OrderType.CMT)
        self.assertEqual(po.required_arrival, date(2026, 9, 1))
        self.assertEqual(po.lines.get().cogs_snapshot, Decimal("50000"))
        children = list(plan.child_purchase_orders.order_by("child_code"))
        self.assertEqual(
            [(child.child_code, child.order_type) for child in children],
            [
                ("A", PurchaseOrder.OrderType.MAIN_MATERIAL),
                ("B", PurchaseOrder.OrderType.CMT),
            ],
        )
        self.assertEqual(children[0].items.get().description, "Fabric")

        self.client.force_login(self.user)
        delete_response = self.client.post(reverse("purchasing:po_delete_draft", args=[po.id]))
        self.assertEqual(delete_response.status_code, 302)
        plan.refresh_from_db()
        self.assertEqual(plan.status, PPICProductionPlan.Status.DRAFT)
        self.assertIsNone(plan.purchase_order_id)

    def test_cmt_fractional_cost_is_saved_as_whole_rupiah_and_generates_po(self):
        incoming = self._approved_incoming(100)
        requirement = PPICRequirement.objects.get(incoming_plan=incoming)
        plan = save_ppic_production_plan(
            actor=self.user,
            requirements=[requirement],
            quantities={str(requirement.id): "100"},
            values={
                "production_method": PPICProductionPlan.ProductionMethod.CMT,
                "supplier": self.supplier,
                "control_material_purchase": False,
                "control_trial": False,
                "control_cut": False,
                "control_additional_process": False,
                "control_make": False,
                "control_trim": False,
                "control_qc": False,
                "control_inbound": True,
                "target_inbound_date": date(2026, 9, 1),
            },
            materials=[{
                "product": self.sku.product_variant.product,
                "supplier": self.supplier,
                "material_name": "Fabric",
                "quantity": Decimal("1"),
                "unit": "meter",
                "unit_price": Decimal("3350"),
            }],
        )

        self.assertEqual(plan.lines.get().unit_cogs, Decimal("34"))
        self.assertEqual(generate_po_from_production_plan(plan.id, self.user).lines.get().cogs_snapshot, Decimal("34"))

    def test_same_material_is_combined_in_supplier_po(self):
        first_requirement = PPICRequirement.objects.get(incoming_plan=self._approved_incoming(100))
        second_product = Product.objects.create(
            code="P-2",
            name="Second Product",
            status=self.sku.product_variant.product.status,
            category=self.sku.product_variant.product.category,
        )
        second_sku = SKU.objects.create(
            sku="SKU-2",
            product_variant=ProductVariant.objects.create(product=second_product, name="Blue"),
            current_retail_price=Decimal("200000"),
            current_master_cogs=Decimal("100000"),
        )
        second_projection = SalesProjection.objects.create(
            scenario=self.scenario,
            month=date(2026, 9, 1),
            sku=second_sku,
            system_recommendation=Decimal("50"),
        )
        approve_sales_projection(second_projection.id, 50, self.user)
        second_incoming = create_incoming_plan(second_projection.id, 0)
        approve_incoming_plan(second_incoming.id, 50, self.user)
        second_requirement = PPICRequirement.objects.get(incoming_plan=second_incoming)
        plan = save_ppic_production_plan(
            actor=self.user,
            requirements=[first_requirement, second_requirement],
            quantities={str(first_requirement.id): "100", str(second_requirement.id): "50"},
            values={
                "production_method": PPICProductionPlan.ProductionMethod.CMT,
                "supplier": self.supplier,
                "control_material_purchase": True,
                "control_trial": False,
                "control_cut": False,
                "control_additional_process": False,
                "control_make": False,
                "control_trim": False,
                "control_qc": False,
                "control_inbound": True,
                "target_material_purchase_date": date(2026, 8, 10),
                "target_inbound_date": date(2026, 9, 1),
            },
            materials=[
                {
                    "product": self.sku.product_variant.product,
                    "supplier": self.supplier,
                    "material_name": "Stripe Blue",
                    "quantity": Decimal("180"),
                    "unit": "Yard",
                    "unit_price": Decimal("19000"),
                },
                {
                    "product": second_product,
                    "supplier": self.supplier,
                    "material_name": "Stripe Blue",
                    "quantity": Decimal("200"),
                    "unit": "Yard",
                    "unit_price": Decimal("19000"),
                },
            ],
        )

        po = generate_po_from_production_plan(plan.id, self.user)
        item = po.items.get()
        self.assertEqual(item.description, "Stripe Blue")
        self.assertEqual(item.quantity, Decimal("380"))
        self.assertIsNone(item.product)
        self.assertEqual(item.total_cost, Decimal("7220000"))
        self.client.force_login(self.user)
        detail = self.client.get(reverse("purchasing:po_detail", args=[po.id]))
        printed = self.client.get(reverse("purchasing:po_print", args=[po.id]))
        self.assertContains(detail, "Product, Second Product")
        self.assertContains(printed, "Product, Second Product")
        self.assertNotContains(detail, "Multiple Products")

    def test_custom_additional_process_is_costed_and_monitored_after_cut(self):
        incoming = self._approved_incoming(100)
        requirement = PPICRequirement.objects.get(incoming_plan=incoming)
        plan = save_ppic_production_plan(
            actor=self.user,
            requirements=[requirement],
            quantities={str(requirement.id): "100"},
            values={
                "production_method": PPICProductionPlan.ProductionMethod.CMT,
                "supplier": self.supplier,
                "control_material_purchase": False,
                "control_trial": False,
                "control_cut": True,
                "control_additional_process": True,
                "control_make": True,
                "control_trim": False,
                "control_qc": False,
                "control_inbound": False,
                "target_cut_start_date": date(2026, 8, 10),
                "target_cut_end_date": date(2026, 8, 11),
                "additional_process_name": "Bordir",
                "target_additional_process_start_date": date(2026, 8, 12),
                "target_additional_process_end_date": date(2026, 8, 13),
                "target_make_start_date": date(2026, 8, 14),
                "target_make_end_date": date(2026, 8, 20),
            },
            materials=[{
                "product": self.sku.product_variant.product,
                "supplier": self.supplier,
                "material_name": "Fabric",
                "quantity": Decimal("100"),
                "unit": "meter",
                "unit_price": Decimal("5000"),
            }, {
                "product": self.sku.product_variant.product,
                "supplier": self.supplier,
                "material_type": PPICProductionMaterial.MaterialType.SUPPORTING,
                "material_name": "Label",
                "quantity": Decimal("100"),
                "unit": "pcs",
                "unit_price": Decimal("0"),
            }],
            process_costs=[{
                "product": self.sku.product_variant.product,
                "cmt_supplier": self.supplier,
                "cmt_unit_cost": Decimal("4000"),
                "additional_process_name": "Bordir",
                "additional_process_supplier": self.supplier,
                "additional_process_unit_cost": Decimal("2000"),
            }],
        )

        self.assertEqual(plan.planned_total_cost, Decimal("1100000"))
        self.assertEqual(plan.planned_unit_cogs, Decimal("11000"))
        po = release_po(generate_po_from_production_plan(plan.id, self.user).id, self.user)
        plan.refresh_from_db()
        children = list(plan.child_purchase_orders.order_by("child_code"))
        self.assertEqual(
            [(child.child_code, child.order_type) for child in children],
            [
                ("A", PurchaseOrder.OrderType.MAIN_MATERIAL),
                ("B", PurchaseOrder.OrderType.CMT),
                ("C", PurchaseOrder.OrderType.ADDITIONAL_PROCESS),
                ("D", PurchaseOrder.OrderType.SUPPORTING_MATERIAL),
            ],
        )
        self.assertTrue(plan.ppo_number.startswith("PPO-VOB-"))
        self.assertEqual(
            [child.po_number[-1] for child in children],
            ["A", "B", "C", "D"],
        )
        self.assertTrue(all(child.status == PurchaseOrder.Status.RELEASED for child in children))
        production_order = po.production_order
        monitoring_plan = production_order.plan
        self.assertEqual(monitoring_plan.additional_process_name, "Bordir")
        self.assertEqual(monitoring_plan.target_additional_process_start_date, date(2026, 8, 12))

        self.client.force_login(self.user)
        detail = self.client.get(reverse("production:detail", args=[production_order.id]))
        self.assertContains(detail, plan.ppo_number)
        self.assertContains(detail, "Bordir")
        self.assertContains(detail, 'data-process-code="ADDITIONAL_PROCESS"')
        self.assertContains(detail, 'data-process-card="ADDITIONAL_PROCESS"')

        choices = dict(eligible_activity_choices(production_order))
        self.assertIn(ProductionActivity.ActivityType.CUT, choices)
        submit_cmt_activity_batch(
            production_order=production_order,
            activity_type=ProductionActivity.ActivityType.CUT,
            activity_date=date(2026, 8, 11),
            line_quantities=((po.lines.get(), 100),),
            actor=self.user,
        )
        choices = dict(eligible_activity_choices(production_order))
        self.assertIn(ProductionActivity.ActivityType.ADDITIONAL_PROCESS, choices)
        self.assertIn("Bordir", choices[ProductionActivity.ActivityType.ADDITIONAL_PROCESS])
        submit_cmt_activity_batch(
            production_order=production_order,
            activity_type=ProductionActivity.ActivityType.ADDITIONAL_PROCESS,
            activity_date=date(2026, 8, 13),
            line_quantities=((po.lines.get(), 100),),
            actor=self.user,
        )
        self.assertIn(
            ProductionActivity.ActivityType.MAKE,
            dict(eligible_activity_choices(production_order)),
        )
        cancel_po(po.id, self.user, "PPO cancellation test")
        self.assertFalse(
            plan.child_purchase_orders.exclude(status=PurchaseOrder.Status.CANCELLED).exists()
        )

    def test_fpp_production_plan_uses_vendor_unit_price(self):
        incoming = self._approved_incoming(100)
        requirement = PPICRequirement.objects.get(incoming_plan=incoming)
        plan = save_ppic_production_plan(
            actor=self.user,
            requirements=[requirement],
            quantities={str(requirement.id): "100"},
            unit_cogs={str(requirement.id): "125,000"},
            values={
                "production_method": PPICProductionPlan.ProductionMethod.FPP,
                "supplier": self.supplier,
                "control_material_purchase": False,
                "control_trial": True,
                "control_cut": False,
                "control_make": True,
                "control_trim": True,
                "control_qc": False,
                "control_inbound": False,
                "target_trial_date": date(2026, 8, 16),
                "target_make_start_date": date(2026, 8, 17),
                "target_make_end_date": date(2026, 8, 25),
                "target_trim_start_date": date(2026, 8, 26),
                "target_trim_end_date": date(2026, 8, 28),
                "notes": "FPP test",
            },
            materials=[],
        )

        self.assertEqual(plan.planned_unit_cogs, Decimal("125000"))
        self.assertEqual(plan.lines.get().unit_cogs, Decimal("125000"))
        po = generate_po_from_production_plan(plan.id, self.user)
        self.assertEqual(po.lines.get().cogs_snapshot, Decimal("125000"))
        self.assertIsNone(po.required_arrival)
        self.assertEqual(po.parent_plan_id, plan.id)
        self.assertEqual(po.child_code, "")
        self.assertEqual(plan.child_purchase_orders.count(), 1)

        po = release_po(po.id, self.user)
        plan.refresh_from_db()
        self.assertEqual(
            po.po_number.removeprefix("PO-"),
            plan.ppo_number.removeprefix("PPO-"),
        )
        self.assertEqual(po.sequence, plan.ppo_sequence)
        monitoring_plan = po.production_order.plan
        self.assertEqual(monitoring_plan.status, ProductionPlan.Status.ACTIVE)
        self.assertEqual(monitoring_plan.target_make_start_date, date(2026, 8, 17))
        self.assertEqual(monitoring_plan.target_trim_end_date, date(2026, 8, 28))

        self.client.force_login(self.user)
        response = self.client.get(reverse("production:monitoring"))
        self.assertContains(response, po.po_number)
        self.assertContains(response, "FPP · Full Package Production")
        monitoring_row = next(
            row for row in response.context["rows"] if row["production_order"].po_id == po.id
        )
        self.assertEqual(monitoring_row["current_code"], "TRIAL")
        self.assertEqual(monitoring_row["next_process_label"], "Trial Production")
        self.assertEqual(monitoring_row["material_status_display"], "Not Controlled")

        detail_response = self.client.get(reverse("production:detail", args=[po.production_order.id]))
        self.assertContains(detail_response, 'data-metric-card="production-method"')
        self.assertContains(detail_response, "FPP · Full Package Production")
        self.assertNotContains(detail_response, 'data-process-card="MATERIAL_PURCHASE"')
        self.assertNotContains(detail_response, 'data-process-card="CUT"')
        self.assertNotContains(detail_response, 'data-process-card="QC"')
        self.assertNotContains(detail_response, 'data-process-card="INBOUND"')
        self.assertContains(detail_response, 'data-process-card="TRIAL"')
        self.assertContains(detail_response, 'data-process-card="MAKE"')
        self.assertContains(detail_response, 'data-process-card="TRIM"')
        self.assertContains(detail_response, 'data-process-code="TRIAL"')
        self.assertContains(detail_response, 'data-process-code="MAKE"')
        self.assertContains(detail_response, 'data-process-code="TRIM"')

        choices = dict(eligible_activity_choices(po.production_order))
        self.assertNotIn(ProductionActivity.ActivityType.MATERIAL_PURCHASE, choices)
        self.assertEqual(
            choices,
            {ProductionActivity.ActivityType.TRIAL_SUBMIT: "Submit Trial Production"},
        )
        submit_production_activity(
            production_order=po.production_order,
            activity_type=ProductionActivity.ActivityType.TRIAL_SUBMIT,
            activity_date=date(2026, 8, 16),
            actor=self.user,
        )
        submit_production_activity(
            production_order=po.production_order,
            activity_type=ProductionActivity.ActivityType.TRIAL_APPROVE,
            activity_date=date(2026, 8, 16),
            actor=self.user,
        )
        choices = dict(eligible_activity_choices(po.production_order))
        self.assertIn(ProductionActivity.ActivityType.MAKE, choices)
        self.assertNotIn(ProductionActivity.ActivityType.CUT, choices)
        submit_cmt_activity_batch(
            production_order=po.production_order,
            activity_type=ProductionActivity.ActivityType.MAKE,
            activity_date=date(2026, 8, 17),
            line_quantities=[(po.lines.get(), 100)],
            actor=self.user,
        )
        self.assertEqual(
            po.production_order.stages.get(stage=ProductionStage.Stage.MAKE).status,
            ProductionStage.Status.COMPLETE,
        )

    def test_requirement_review_opens_production_plan_without_creating_po(self):
        incoming = self._approved_incoming(100)
        requirement = PPICRequirement.objects.get(incoming_plan=incoming)
        second_product = Product.objects.create(
            code="P-2",
            name="Second Product",
            status=self.sku.product_variant.product.status,
            category=self.sku.product_variant.product.category,
        )
        second_sku = SKU.objects.create(
            sku="SKU-2",
            product_variant=ProductVariant.objects.create(product=second_product, name="Navy"),
            current_retail_price=Decimal("200000"),
            current_master_cogs=Decimal("100000"),
        )
        second_projection = SalesProjection.objects.create(
            scenario=self.scenario,
            month=date(2026, 9, 1),
            sku=second_sku,
            system_recommendation=Decimal("50"),
        )
        approve_sales_projection(second_projection.id, 50, self.user)
        second_incoming = create_incoming_plan(second_projection.id, 0)
        approve_incoming_plan(second_incoming.id, 50, self.user)
        second_requirement = PPICRequirement.objects.get(incoming_plan=second_incoming)
        third_product = Product.objects.create(
            code="P-3",
            name="Addable Product",
            status=self.sku.product_variant.product.status,
            category=self.sku.product_variant.product.category,
        )
        third_sku = SKU.objects.create(
            sku="SKU-3",
            product_variant=ProductVariant.objects.create(product=third_product, name="Black"),
            current_retail_price=Decimal("150000"),
            current_master_cogs=Decimal("75000"),
        )
        third_projection = SalesProjection.objects.create(
            scenario=self.scenario,
            month=date(2026, 9, 1),
            sku=third_sku,
            system_recommendation=Decimal("30"),
        )
        approve_sales_projection(third_projection.id, 30, self.user)
        third_incoming = create_incoming_plan(third_projection.id, 0)
        approve_incoming_plan(third_incoming.id, 30, self.user)
        third_requirement = PPICRequirement.objects.get(incoming_plan=third_incoming)
        self.client.force_login(self.user)

        review_response = self.client.post(
            reverse("purchasing:production_plan_create"),
            {"action": "review", "requirement": [str(requirement.id), str(second_requirement.id)]},
        )
        self.assertRedirects(review_response, reverse("purchasing:production_plan_create"))
        response = self.client.get(reverse("purchasing:production_plan_create"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Buat Production Plan")
        self.assertContains(response, "CMT · Cut, Make, Trim")
        self.assertContains(response, "COGS / Qty")
        self.assertContains(response, "Total COGS")
        self.assertContains(response, "2 Products Selected")
        self.assertContains(response, "Product Name")
        self.assertContains(response, "Product / Parent SKU")
        self.assertContains(response, 'name="material_product"')
        self.assertContains(response, "CMT Cost / Qty")
        self.assertContains(response, "Biaya Additional Process per Product")
        self.assertContains(response, "Vendor CMT")
        self.assertContains(response, "Supplier")
        self.assertContains(response, "Material Type")
        self.assertContains(response, "Supporting Material")
        self.assertContains(response, "Bordir / Sablon")
        self.assertContains(response, 'name="cmt_cost_product"')
        self.assertContains(response, 'name="additional_cost_product"')
        self.assertContains(response, "data-remove-requirement")
        self.assertContains(response, "+ Add Product")
        self.assertContains(response, 'data-add-product-template')
        self.assertContains(response, "Addable Product")
        self.assertContains(response, "Second Product")
        self.assertContains(response, "data-draft-storage-key")
        self.assertContains(response, "sessionStorage")
        self.assertNotContains(response, ">Variant<")
        self.assertNotContains(response, ">Generate Draft PO</button>")
        self.assertEqual(PPICProductionPlan.objects.count(), 0)
        self.assertEqual(PurchaseOrder.objects.count(), 0)

        plan = save_ppic_production_plan(
            actor=self.user,
            requirements=[requirement, second_requirement],
            quantities={str(requirement.id): "100", str(second_requirement.id): "50"},
            values={
                "production_method": PPICProductionPlan.ProductionMethod.CMT,
                "supplier": self.supplier,
            },
            materials=[
                {
                    "product": self.sku.product_variant.product,
                    "material_name": "Main fabric",
                    "quantity": Decimal("100"),
                    "unit": "meter",
                    "unit_price": Decimal("5000"),
                },
                {
                    "product": second_product,
                    "material_name": "Second fabric",
                    "quantity": Decimal("50"),
                    "unit": "meter",
                    "unit_price": Decimal("6000"),
                },
            ],
            process_costs=[
                {
                    "product": self.sku.product_variant.product,
                    "cmt_unit_cost": Decimal("1000"),
                    "additional_process_unit_cost": Decimal("0"),
                },
                {
                    "product": second_product,
                    "cmt_unit_cost": Decimal("2000"),
                    "additional_process_unit_cost": Decimal("0"),
                },
            ],
        )
        self.assertEqual(plan.lines.count(), 2)
        self.assertSetEqual(
            set(plan.materials.values_list("product_id", flat=True)),
            {self.sku.product_variant.product_id, second_product.id},
        )
        self.assertEqual(plan.lines.get(requirement=requirement).unit_cogs, Decimal("6000.0000"))
        self.assertEqual(plan.lines.get(requirement=second_requirement).unit_cogs, Decimal("8000.0000"))

        response = self.client.post(
            reverse("purchasing:production_plan_detail", args=[plan.id]),
            {
                "action": "save",
                "requirement": [str(requirement.id), str(third_requirement.id)],
                "production_method": PPICProductionPlan.ProductionMethod.CMT,
                "supplier": str(self.supplier.id),
                "control_material_purchase": "on",
                "control_additional_process": "on",
                "additional_process_name": "Bordir",
                f"quantity_{requirement.id}": "100",
                f"unit_cogs_{requirement.id}": "0",
                f"quantity_{third_requirement.id}": "30",
                f"unit_cogs_{third_requirement.id}": "0",
                "material_product": [str(self.sku.product_variant.product_id)],
                "material_name": ["Main fabric"],
                "material_quantity": ["100"],
                "material_unit": ["meter"],
                "material_unit_price": ["5000"],
                "material_supplier": [str(self.supplier.id)],
                "material_type": ["SUPPORTING"],
                "cmt_cost_product": [
                    str(self.sku.product_variant.product_id),
                    str(third_product.id),
                ],
                "cmt_unit_cost": ["1,000", "2,000"],
                "cmt_supplier": [str(self.supplier.id), str(self.supplier.id)],
                "additional_cost_product": [str(self.sku.product_variant.product_id)],
                "additional_process_cost_name": ["Bordir"],
                "additional_process_supplier": [str(self.supplier.id)],
                "additional_process_unit_cost": ["500"],
            },
        )
        self.assertRedirects(response, reverse("purchasing:production_plan_detail", args=[plan.id]))
        self.assertEqual(plan.lines.count(), 2)
        self.assertFalse(plan.lines.filter(requirement=second_requirement).exists())
        self.assertEqual(plan.lines.get(requirement=third_requirement).quantity, Decimal("30"))
        self.assertEqual(
            plan.lines.get(requirement=requirement).unit_cogs,
            Decimal("6500.0000"),
        )
        self.assertEqual(
            plan.lines.get(requirement=third_requirement).unit_cogs,
            Decimal("2000.0000"),
        )
        self.assertEqual(plan.materials.get().supplier, self.supplier)
        self.assertEqual(plan.materials.get().material_type, "SUPPORTING")
        first_product_cost = plan.process_costs.get(product=self.sku.product_variant.product)
        self.assertEqual(first_product_cost.cmt_supplier, self.supplier)
        self.assertEqual(first_product_cost.additional_process_name, "Bordir")
        self.assertEqual(first_product_cost.additional_process_supplier, self.supplier)
        requirement_response = self.client.get(reverse("purchasing:requirements"))
        returned_requirement = next(
            row for row in requirement_response.context["requirements"] if row.id == second_requirement.id
        )
        self.assertEqual(returned_requirement.available_for_plan_qty, Decimal("50"))

    def test_fpp_draft_can_be_saved_before_cogs_and_process_dates_are_complete(self):
        incoming = self._approved_incoming(100)
        requirement = PPICRequirement.objects.get(incoming_plan=incoming)
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("purchasing:production_plan_create"),
            {
                "action": "save",
                "requirement": str(requirement.id),
                "production_method": PPICProductionPlan.ProductionMethod.FPP,
                "control_make": "on",
                "control_trim": "on",
                f"quantity_{requirement.id}": "100",
                f"unit_cogs_{requirement.id}": "",
            },
        )

        plan = PPICProductionPlan.objects.get()
        self.assertRedirects(response, reverse("purchasing:production_plan_detail", args=[plan.id]))
        self.assertEqual(plan.need_month, requirement.need_month)
        self.assertEqual(plan.lines.get().unit_cogs, Decimal("0"))
        with self.assertRaisesMessage(ValidationError, "COGS per quantity FPP wajib diisi"):
            generate_po_from_production_plan(plan.id, self.user)

    def test_failed_generate_keeps_submitted_process_dates(self):
        incoming = self._approved_incoming(100)
        requirement = PPICRequirement.objects.get(incoming_plan=incoming)
        plan = save_ppic_production_plan(
            actor=self.user,
            requirements=[requirement],
            quantities={str(requirement.id): "100"},
            unit_cogs={str(requirement.id): "125000"},
            values={
                "production_method": PPICProductionPlan.ProductionMethod.FPP,
                "supplier": self.supplier,
                "control_material_purchase": False,
                "control_trial": True,
                "control_cut": False,
                "control_make": True,
                "control_trim": True,
                "control_qc": True,
                "control_inbound": True,
            },
            materials=[],
        )
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("purchasing:production_plan_detail", args=[plan.id]),
            {
                "action": "generate",
                "requirement": [str(requirement.id)],
                "production_method": PPICProductionPlan.ProductionMethod.FPP,
                "supplier": str(self.supplier.id),
                "control_trial": "on",
                "control_make": "on",
                "control_trim": "on",
                "control_qc": "on",
                "control_inbound": "on",
                "target_make_start_date": "2026-08-17",
                "target_make_end_date": "2026-08-25",
                "target_trim_start_date": "2026-08-26",
                "target_trim_end_date": "2026-08-28",
                "target_qc_start_date": "2026-08-29",
                "target_qc_end_date": "2026-08-30",
                "target_inbound_date": "2026-09-01",
                f"quantity_{requirement.id}": "100",
                f"unit_cogs_{requirement.id}": "125000",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Target Trial Production wajib diisi.")
        self.assertContains(response, 'value="2026-08-17"')
        plan.refresh_from_db()
        self.assertEqual(plan.status, PPICProductionPlan.Status.DRAFT)
        self.assertEqual(plan.target_make_start_date, date(2026, 8, 17))
        self.assertEqual(plan.target_trim_end_date, date(2026, 8, 28))
        self.assertEqual(plan.target_qc_end_date, date(2026, 8, 30))
        self.assertEqual(plan.target_inbound_date, date(2026, 9, 1))
        self.assertIsNone(plan.purchase_order_id)

    def test_po_generator_lists_saved_draft_plans_and_honors_month_filter(self):
        first_plan = self._approved_incoming(100)
        first = PPICRequirement.objects.get(incoming_plan=first_plan)
        second_scenario = ProjectionScenario.objects.create(
            name="October 2026",
            start_month=date(2026, 10, 1),
            end_month=date(2026, 10, 1),
            created_by=self.user,
        )
        projection = SalesProjection.objects.create(
            scenario=second_scenario,
            month=date(2026, 10, 1),
            sku=self.sku,
            system_recommendation=Decimal("40"),
        )
        approve_sales_projection(projection.id, 40, self.user)
        second_plan = create_incoming_plan(projection.id, 0)
        approve_incoming_plan(second_plan.id, 40, self.user)
        second = PPICRequirement.objects.get(incoming_plan=second_plan)

        september_production_plan = save_ppic_production_plan(
            actor=self.user,
            requirements=[first],
            quantities={str(first.id): "100"},
            unit_cogs={str(first.id): "115000"},
            values={
                "production_method": PPICProductionPlan.ProductionMethod.FPP,
                "supplier": self.supplier,
                "control_material_purchase": False,
                "control_trial": False,
                "control_cut": False,
                "control_make": True,
                "control_trim": False,
                "control_qc": False,
                "control_inbound": False,
                "target_make_start_date": date(2026, 8, 17),
                "target_make_end_date": date(2026, 8, 25),
            },
            materials=[],
        )
        october_production_plan = save_ppic_production_plan(
            actor=self.user,
            requirements=[second],
            quantities={str(second.id): "40"},
            unit_cogs={str(second.id): "125000"},
            values={
                "production_method": PPICProductionPlan.ProductionMethod.FPP,
                "supplier": self.supplier,
                "control_material_purchase": False,
                "control_trial": False,
                "control_cut": False,
                "control_make": True,
                "control_trim": False,
                "control_qc": False,
                "control_inbound": False,
                "target_make_start_date": date(2026, 9, 17),
                "target_make_end_date": date(2026, 9, 25),
            },
            materials=[],
        )
        self.client.force_login(self.user)
        response = self.client.get(reverse("purchasing:generator"), {"month": "2026-09"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["plans"], [september_production_plan])
        self.assertEqual(response.context["selected_month"], "2026-09")
        self.assertContains(response, "Review &amp; Generate PO")
        self.assertEqual(PurchaseOrder.objects.count(), 0)

        po = generate_po_from_production_plan(september_production_plan.id, self.user)
        all_plans = self.client.get(reverse("purchasing:generator"))
        self.assertEqual(all_plans.context["plans"], [october_production_plan])

        self.client.post(reverse("purchasing:po_delete_draft", args=[po.id]))
        restored = self.client.get(reverse("purchasing:generator"), {"month": "2026-09"})
        self.assertEqual(restored.context["plans"], [september_production_plan])

    def test_tracking_page_is_available_without_test_sheet_data(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("purchasing:tracking"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Belum ada PO aktual")

    def test_ppic_navigation_includes_pre_po_production_plan(self):
        self.client.force_login(self.user)
        session = self.client.session
        session["active_module"] = "operation"
        session.save()
        expected = {
            "purchasing:requirements": "PPIC Requirement",
            "purchasing:production_plans": "Production Plan",
            "purchasing:generator": "Create dan Print Purchase Order",
            "purchasing:purchase_orders": "Daftar dan Tracking Purchase Order",
            "purchasing:vendors": "Daftar Vendor",
        }
        for route, heading in expected.items():
            with self.subTest(route=route):
                response = self.client.get(reverse(route))
                self.assertEqual(response.status_code, 200)
                self.assertContains(response, heading)
                self.assertContains(response, "Requirement")
                self.assertContains(response, "Production Plan")
                self.assertContains(response, "PO Generator")
                self.assertContains(response, "Purchase Order")
                self.assertContains(response, "Daftar Vendor")
                for group_name in ("merchandising", "purchasing", "production", "inventory"):
                    self.assertContains(response, f'data-nav-group-toggle="{group_name}"')
                    self.assertContains(response, f'data-nav-group-panel="{group_name}"')
                self.assertContains(response, "Planning Builder")
                self.assertContains(response, "Production Monitoring")
                self.assertContains(response, "Inventory Summary")
                self.assertNotContains(response, "Manual Requirement Product")
                if route == "purchasing:generator":
                    self.assertNotContains(response, "PO produk baru")

    def test_requirement_need_month_filter_updates_table_and_kpis(self):
        september_plan = self._approved_incoming(100)
        october_scenario = ProjectionScenario.objects.create(
            name="October 2026",
            start_month=date(2026, 10, 1),
            end_month=date(2026, 10, 1),
            created_by=self.user,
        )
        october_projection = SalesProjection.objects.create(
            scenario=october_scenario,
            month=date(2026, 10, 1),
            sku=self.sku,
            system_recommendation=Decimal("40"),
        )
        approve_sales_projection(october_projection.id, 40, self.user)
        october_plan = create_incoming_plan(october_projection.id, 0)
        approve_incoming_plan(october_plan.id, 40, self.user)

        self.client.force_login(self.user)
        response = self.client.get(reverse("purchasing:requirements"), {"need_month": "2026-09"})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["selected_need_month"], "2026-09")
        self.assertEqual(response.context["total_approved"], Decimal("100"))
        self.assertEqual(response.context["total_remaining"], Decimal("100"))
        self.assertEqual(len(response.context["requirements"]), 1)
        self.assertContains(response, "September 2026")
        self.assertContains(response, "All Months")
        self.assertEqual(
            {month.strftime("%Y-%m") for month in response.context["need_month_options"]},
            {september_plan.month.strftime("%Y-%m"), october_plan.month.strftime("%Y-%m")},
        )

    def test_requirement_status_and_category_filters_are_cascading(self):
        first_plan = self._approved_incoming(100)
        second_status = ProductStatus.objects.create(code="ESSENTIAL", name="Essential+")
        second_category = Category.objects.create(code="KNITWEAR", name="Knitwear")
        second_product = Product.objects.create(
            code="P-2",
            name="Second Product",
            status=second_status,
            category=second_category,
        )
        second_variant = ProductVariant.objects.create(product=second_product, name="Blue", color="Blue")
        second_sku = SKU.objects.create(
            sku="SKU-2",
            product_variant=second_variant,
            current_retail_price=Decimal("250000"),
            current_master_cogs=Decimal("110000"),
        )
        second_projection = SalesProjection.objects.create(
            scenario=self.scenario,
            month=date(2026, 9, 1),
            sku=second_sku,
            system_recommendation=Decimal("30"),
        )
        approve_sales_projection(second_projection.id, 30, self.user)
        second_plan = create_incoming_plan(second_projection.id, 0)
        approve_incoming_plan(second_plan.id, 30, self.user)

        self.client.force_login(self.user)
        response = self.client.get(
            reverse("purchasing:requirements"),
            {
                "need_month": "2026-09",
                "status": str(second_status.id),
                "category": str(second_category.id),
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["selected_status"], str(second_status.id))
        self.assertEqual(response.context["selected_category"], str(second_category.id))
        self.assertEqual(response.context["total_approved"], Decimal("30"))
        self.assertEqual(len(response.context["requirements"]), 1)
        self.assertEqual(response.context["requirements"][0].sku_id, second_sku.id)
        self.assertEqual(
            response.context["category_options"],
            [{"value": str(second_category.id), "label": "Knitwear"}],
        )

        stale_response = self.client.get(
            reverse("purchasing:requirements"),
            {
                "need_month": "2026-09",
                "status": str(self.sku.product_variant.product.status_id),
                "category": str(second_category.id),
            },
        )
        self.assertEqual(stale_response.context["selected_category"], "")
        self.assertEqual(
            stale_response.context["category_options"],
            [{"value": str(self.sku.product_variant.product.category_id), "label": "Apparel"}],
        )

        first_requirement = PPICRequirement.objects.get(incoming_plan=first_plan)
        create_draft_po(
            supplier=self.supplier,
            need_month=date(2026, 9, 1),
            actor=self.user,
            requirement_quantities={first_requirement.id: 100},
        )
        allocation_response = self.client.get(
            reverse("purchasing:requirements"),
            {
                "need_month": "2026-09",
                "allocation_status": ["Open", "Fully Allocated"],
            },
        )
        self.assertEqual(
            allocation_response.context["selected_allocation_statuses"],
            ["Open", "Fully Allocated"],
        )
        self.assertEqual(
            allocation_response.context["allocation_status_options"],
            [
                {"value": "Open", "label": "Open"},
                {"value": "Fully Allocated", "label": "Fully Allocated"},
            ],
        )
        self.assertEqual(len(allocation_response.context["requirements"]), 2)
        fully_allocated_only = self.client.get(
            reverse("purchasing:requirements"),
            {"need_month": "2026-09", "allocation_status": "Fully Allocated"},
        )
        self.assertEqual(fully_allocated_only.context["total_approved"], Decimal("130"))
        self.assertEqual(fully_allocated_only.context["total_ordered"], Decimal("100"))
        self.assertEqual(fully_allocated_only.context["total_remaining"], Decimal("30"))
        self.assertEqual(fully_allocated_only.context["requirement_line_count"], 2)
        self.assertEqual(len(fully_allocated_only.context["requirements"]), 1)
        self.assertEqual(fully_allocated_only.context["requirements"][0].sku_id, self.sku.id)

    def test_purchase_order_list_can_search_and_filter(self):
        product = self.sku.product_variant.product
        po = create_draft_po(
            supplier=self.supplier,
            need_month=date(2026, 9, 1),
            actor=self.user,
            manual_lines=[(self.sku, 10)],
        )
        po = release_po(po.id, self.user)
        second_product = Product.objects.create(
            code="P-PO-2",
            name="Other Product",
            status=product.status,
            category=product.category,
        )
        second_sku = SKU.objects.create(
            sku="SKU-PO-2",
            product_variant=ProductVariant.objects.create(product=second_product, name="Navy"),
            current_master_cogs=Decimal("90000"),
        )
        create_draft_po(
            supplier=self.supplier,
            need_month=date(2026, 9, 1),
            actor=self.user,
            manual_lines=[(second_sku, 5)],
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse("purchasing:purchase_orders"), {"q": po.po_number})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context["rows"]), 1)
        self.assertEqual(response.context["rows"][0]["line_count"], 1)
        self.assertContains(response, po.po_number)

        response = self.client.get(reverse("purchasing:purchase_orders"), {"q": "NOT-FOUND"})
        self.assertEqual(len(response.context["rows"]), 0)

        response = self.client.get(
            reverse("purchasing:purchase_orders"),
            {"product": str(product.id)},
        )
        self.assertEqual(len(response.context["rows"]), 1)
        self.assertEqual(response.context["rows"][0]["po"].id, po.id)
        self.assertEqual(response.context["selected"]["products"], [str(product.id)])
        self.assertEqual(
            {item.id for item in response.context["product_options"]},
            {product.id, second_product.id},
        )

        response = self.client.get(
            reverse("purchasing:purchase_orders"),
            {"product": [str(product.id), str(second_product.id)]},
        )
        self.assertEqual(len(response.context["rows"]), 2)

        response = self.client.get(
            reverse("purchasing:purchase_orders"),
            {"supplier": str(self.supplier.id), "need_month": "2026-09", "po_status": "RELEASED", "schedule_status": "Open"},
        )
        self.assertEqual(len(response.context["rows"]), 1)

    def test_purchase_order_list_groups_sku_lines_and_detail_expands_them(self):
        second_sku = SKU.objects.create(
            sku="SKU-2",
            product_variant=self.sku.product_variant,
            current_retail_price=Decimal("250000"),
            current_master_cogs=Decimal("125000"),
        )
        po = create_draft_po(
            supplier=self.supplier,
            need_month=date(2026, 9, 1),
            actor=self.user,
            manual_lines=[(self.sku, 10), (second_sku, 20)],
        )
        po = release_po(po.id, self.user)
        first_line = po.lines.order_by("created_at").first()
        first_line.qc_passed_before_cutover_qty = Decimal("10")
        first_line.save(update_fields=["qc_passed_before_cutover_qty"])
        self.client.force_login(self.user)

        response = self.client.get(reverse("purchasing:purchase_orders"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context["rows"]), 1)
        self.assertEqual(response.context["rows"][0]["line_count"], 2)
        self.assertEqual(response.context["rows"][0]["total_qty"], Decimal("30"))
        self.assertIsNone(response.context["rows"][0]["qc_passed"])
        self.assertContains(response, "Jumlah SKU")
        self.assertNotContains(response, "SKU-1")
        self.assertNotContains(response, "SKU-2")

        detail_response = self.client.get(reverse("purchasing:po_detail", args=[po.id]))
        self.assertEqual(detail_response.status_code, 200)
        self.assertContains(detail_response, "SKU-1")
        self.assertContains(detail_response, "SKU-2")

        print_response = self.client.get(reverse("purchasing:po_print", args=[po.id]))
        self.assertEqual(print_response.status_code, 200)
        self.assertContains(print_response, "<h1>Purchase Order</h1>", html=True)
        self.assertContains(print_response, "PO Number")
        self.assertNotContains(print_response, "Need Month")
        self.assertNotContains(print_response, "Purchased By")
        self.assertNotContains(print_response, "Vobia Warehouse")
        self.assertNotContains(print_response, "Document Control")
        self.assertNotContains(print_response, "Notes")
        self.assertContains(print_response, "SKU-1")
        self.assertContains(print_response, "SKU-2")
        self.assertContains(print_response, "Rp 3.500.000")
        self.assertContains(print_response, "30 pcs")
        self.assertContains(print_response, "Created By")
        self.assertContains(print_response, "Checked By")
        self.assertContains(print_response, "Approved By")
        self.assertContains(print_response, "Create, Check &amp; Approve by Digital Sign")
        self.assertNotContains(print_response, "Payment Terms")

        QCInspection.objects.create(
            po_line=first_line,
            inspected_at=timezone.now(),
            qty_inspected=Decimal("3"),
            qty_passed=Decimal("2"),
            qty_failed=Decimal("1"),
            failed_disposition=QCInspection.Disposition.REWORK,
            recorded_by=self.user,
        )
        tracked_after_qc = self.client.get(reverse("purchasing:purchase_orders"))
        self.assertEqual(tracked_after_qc.context["rows"][0]["qc_passed"], Decimal("2"))

    def test_po_detail_shows_approved_final_quantity_and_cogs(self):
        po = release_po(
            create_draft_po(
                supplier=self.supplier,
                need_month=date(2026, 9, 1),
                actor=self.user,
                manual_lines=[(self.sku, 10)],
            ).id,
            self.user,
        )
        line = po.lines.get()
        ProductionCogsFinalization.objects.create(
            production_order=ensure_production_order(po, actor=self.user),
            line_snapshot=[{
                "po_line_id": str(line.id),
                "sellable_qty": "9",
                "final_unit_cogs": "111111.1111",
                "unit_cogs_increase": "11111.1111",
            }],
            total_po_cost=Decimal("1000000"),
            total_final_cost=Decimal("1000000"),
            approved_by=self.user,
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse("purchasing:po_detail", args=[po.id]))

        self.assertContains(response, "Final Qty")
        self.assertContains(response, "Rp 111.111")
        self.assertContains(response, "+Rp 11.111")
        self.assertContains(response, "Total COGS")
        self.assertContains(response, "Rp 1.000.000")
        self.assertNotContains(response, "QC records")
        self.assertContains(response, "Riwayat Quantity Activity")
        self.assertEqual(
            response.context["activity_rows"][0],
            {"label": "Cutting", "qty": Decimal("0"), "po_qty": Decimal("10"), "gap": Decimal("-10")},
        )

    def test_legacy_wip_vendor_revision_is_scoped_and_audited(self):
        corrected_supplier = Supplier.objects.create(code="HARMONI", name="Harmoni")
        po = PurchaseOrder.objects.create(
            po_number="PO-VOB-06/26-087",
            supplier=self.supplier,
            need_month=date(2026, 6, 1),
            source=PurchaseOrder.Source.LEGACY_WIP,
            status=PurchaseOrder.Status.RELEASED,
            created_by=self.user,
            released_by=self.user,
            migration_cutoff_date=date(2026, 7, 31),
            migration_evidence_reference="PO WIP.xlsx · test checksum",
        )
        po = revise_legacy_wip_supplier(
            po.id,
            corrected_supplier,
            self.user,
            "Koreksi berdasarkan PO WIP (1).numbers",
        )
        self.assertEqual(po.supplier, corrected_supplier)
        event = AuditEvent.objects.get(action="legacy_wip_supplier_revised", entity_id=str(po.id))
        self.assertEqual(event.before_values["supplier_name"], "Supplier 1")
        self.assertEqual(event.after_values["supplier_name"], "Harmoni")
        self.assertFalse(event.metadata["quantities_and_cost_snapshots_changed"])

        regular_po = create_draft_po(
            supplier=self.supplier,
            need_month=date(2026, 9, 1),
            actor=self.user,
            manual_lines=[(self.sku, 1)],
        )
        regular_po = release_po(regular_po.id, self.user)
        with self.assertRaises(ValidationError):
            revise_legacy_wip_supplier(
                regular_po.id,
                corrected_supplier,
                self.user,
                "Tidak boleh untuk PO reguler",
            )

    def test_supplemental_wip_requires_evidence_and_allows_vendor_revision(self):
        corrected_supplier = Supplier.objects.create(code="MORPH", name="Morph")
        po = PurchaseOrder(
            po_number="PO-VOB-08/26-124",
            issue_month=date(2026, 8, 1),
            supplier=self.supplier,
            need_month=date(2026, 8, 1),
            source=PurchaseOrder.Source.SUPPLEMENTAL_WIP,
            status=PurchaseOrder.Status.RELEASED,
            created_by=self.user,
            released_by=self.user,
            migration_evidence_reference="PO WIP tambahan.xlsx · test checksum",
        )
        po.full_clean()
        po.save()

        revised = revise_legacy_wip_supplier(
            po.id,
            corrected_supplier,
            self.user,
            "Koreksi vendor berdasarkan file tambahan.",
        )

        self.assertEqual(revised.supplier, corrected_supplier)
        po.migration_cutoff_date = date(2026, 7, 31)
        with self.assertRaises(ValidationError):
            po.full_clean()

    def test_legacy_wip_vendor_revision_view_requires_reason(self):
        corrected_supplier = Supplier.objects.create(code="HARMONI", name="Harmoni")
        po = PurchaseOrder.objects.create(
            po_number="PO-VOB-06/26-087",
            supplier=self.supplier,
            need_month=date(2026, 6, 1),
            source=PurchaseOrder.Source.LEGACY_WIP,
            status=PurchaseOrder.Status.RELEASED,
            created_by=self.user,
            released_by=self.user,
            migration_cutoff_date=date(2026, 7, 31),
            migration_evidence_reference="PO WIP.xlsx · test checksum",
        )
        self.client.force_login(self.user)
        invalid = self.client.post(
            reverse("purchasing:po_revise_vendor", args=[po.id]),
            {"supplier": str(corrected_supplier.id), "reason": ""},
        )
        self.assertEqual(invalid.status_code, 302)
        po.refresh_from_db()
        self.assertEqual(po.supplier, self.supplier)

        valid = self.client.post(
            reverse("purchasing:po_revise_vendor", args=[po.id]),
            {
                "supplier": str(corrected_supplier.id),
                "reason": "Koreksi berdasarkan PO WIP (1).numbers",
            },
        )
        self.assertEqual(valid.status_code, 302)
        po.refresh_from_db()
        self.assertEqual(po.supplier, corrected_supplier)

    def test_unused_vendor_can_be_deleted_but_used_vendor_is_protected(self):
        unused = Supplier.objects.create(code="UNUSED", name="Unused Vendor")
        unused_id = unused.id
        deleted = delete_unused_supplier(
            unused.id,
            self.user,
            "Placeholder vendor tidak lagi digunakan",
        )
        self.assertEqual(deleted["name"], "Unused Vendor")
        self.assertFalse(Supplier.objects.filter(pk=unused_id).exists())
        event = AuditEvent.objects.get(action="unused_supplier_deleted", entity_id=str(unused_id))
        self.assertTrue(event.after_values["deleted"])

        po = create_draft_po(
            supplier=self.supplier,
            need_month=date(2026, 9, 1),
            actor=self.user,
            manual_lines=[(self.sku, 1)],
        )
        with self.assertRaises(ValidationError):
            delete_unused_supplier(
                self.supplier.id,
                self.user,
                "Vendor masih dipakai PO sehingga harus ditolak",
            )
        self.assertTrue(Supplier.objects.filter(pk=self.supplier.id).exists())
        self.assertTrue(PurchaseOrder.objects.filter(pk=po.id).exists())

    def test_po_wip_import_stages_then_commits_without_stock_posting(self):
        source = (
            "NO PO,SKU Induk,SKU,Nama Barang,WIP\n"
            "PO-VOB-06/26-087,,SKU-1,Legacy Product Name,10\n"
        )
        uploaded = SimpleUploadedFile("PO WIP.csv", source.encode(), content_type="text/csv")
        with tempfile.TemporaryDirectory() as directory, override_settings(PRIVATE_UPLOAD_ROOT=directory):
            batch = create_po_wip_import(uploaded, self.user)
            self.assertEqual(batch.status, POWIPImportBatch.Status.READY)
            self.assertEqual(batch.po_count, 1)
            self.assertEqual(batch.total_rows, 1)
            self.assertEqual(batch.total_outstanding_qty, Decimal("10"))
            self.assertEqual(batch.blocking_issue_count, 0)
            self.assertEqual(batch.warning_count, 3)
            approve_po_wip_import(batch.id, self.user)

        po = PurchaseOrder.objects.get(po_number="PO-VOB-06/26-087")
        self.assertEqual(po.source, PurchaseOrder.Source.LEGACY_WIP)
        self.assertEqual(po.status, PurchaseOrder.Status.RELEASED)
        self.assertEqual(po.supplier.name, "Vobia Vendor")
        self.assertEqual(po.need_month, date(2026, 6, 1))
        self.assertEqual(po.created_at.date(), date(2026, 7, 31))
        self.assertIsNone(po.required_arrival)
        line = po.lines.get()
        self.assertEqual(line.ordered_qty, Decimal("10"))
        self.assertEqual(line.received_before_cutover_qty, Decimal("0"))
        self.assertEqual(line.qc_passed_before_cutover_qty, Decimal("10"))
        self.assertEqual(line.cogs_snapshot, Decimal("100000"))
        self.assertEqual(InventoryMovement.objects.count(), 0)
        self.assertEqual(FIFOLayer.objects.count(), 0)

    def test_po_wip_import_blocks_duplicate_po_and_sku(self):
        source = (
            "NO PO,SKU Induk,SKU,Nama Barang,WIP\n"
            "PO-VOB-06/26-087,,SKU-1,Product,10\n"
            "PO-VOB-06/26-087,,SKU-1,Product,5\n"
        )
        uploaded = SimpleUploadedFile("PO WIP duplicate.csv", source.encode(), content_type="text/csv")
        with tempfile.TemporaryDirectory() as directory, override_settings(PRIVATE_UPLOAD_ROOT=directory):
            batch = create_po_wip_import(uploaded, self.user)
        self.assertEqual(batch.status, POWIPImportBatch.Status.BLOCKED)
        self.assertGreater(batch.blocking_issue_count, 0)
        self.assertEqual(batch.issues.filter(code="DUPLICATE_PO_SKU").count(), 2)
