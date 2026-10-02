from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.conf import settings
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.paginator import Paginator
from django.db import transaction
from django.http import HttpResponse, HttpResponseNotAllowed
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from accounts.access import module_level
from sales.views import _save_store_traffic
from traffic.models import StoreTrafficMetric
from .forms import MasterImportUploadForm, SalesImportUploadForm
from .models import (
    ImportValidationIssue,
    MasterImportBatch,
    SalesImportBatch,
    SalesImportIssue,
    StagedSalesRow,
)
from .services.master_commit import approve_master_import, cancel_master_import
from .services.sales_commit import approve_sales_import
from .services.sales_parser import parse_sales_batch
from .services.storage import DuplicateRawFile, create_master_import, create_sales_import
from sales.services.requirements import import_requirements, summarize_import_requirements
from inventory.models import FIFOOpeningImportBatch


def _can_edit_store_traffic(user):
    return user.is_superuser or module_level(user, "sales") in {"edit", "approve"}


@login_required
def master_import_list(request):
    batches = MasterImportBatch.objects.select_related("raw_file", "raw_file__uploaded_by")[:30]
    return render(request, "imports/master_list.html", {"batches": batches})


@login_required
def master_import_upload(request):
    if request.method == "POST":
        form = MasterImportUploadForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                batch = create_master_import(form.cleaned_data["file"], request.user)
            except DuplicateRawFile as exc:
                existing_batch = exc.raw_file.master_batches.order_by("-created_at").first()
                detail = f" Batch sebelumnya: {existing_batch.id}." if existing_batch else ""
                form.add_error("file", "File identik sudah pernah diunggah." + detail)
            else:
                if batch.status == MasterImportBatch.Status.READY:
                    messages.success(request, "File berhasil diparsing dan siap direview.")
                else:
                    messages.warning(request, "File diparsing tetapi memiliki blocking issue.")
                return redirect("imports:master_detail", batch_id=batch.id)
    elif request.method == "GET":
        form = MasterImportUploadForm()
    else:
        return HttpResponseNotAllowed(["GET", "POST"])
    return render(request, "imports/master_upload.html", {"form": form})


@login_required
def master_import_detail(request, batch_id):
    batch = get_object_or_404(
        MasterImportBatch.objects.select_related("raw_file", "approved_by"),
        pk=batch_id,
    )
    action_filter = request.GET.get("action", "")
    severity_filter = request.GET.get("severity", "")

    staged_rows = batch.staged_rows.select_related("existing_sku")
    if action_filter in {choice for choice, _ in batch.staged_rows.model.ProposedAction.choices}:
        staged_rows = staged_rows.filter(proposed_action=action_filter)
    page = Paginator(staged_rows, 50).get_page(request.GET.get("page"))

    issues = batch.issues.select_related("staged_row")
    if severity_filter in {choice for choice, _ in ImportValidationIssue.Severity.choices}:
        issues = issues.filter(severity=severity_filter)

    context = {
        "batch": batch,
        "page": page,
        "issues": issues[:100],
        "action_filter": action_filter,
        "severity_filter": severity_filter,
    }
    return render(request, "imports/master_detail.html", context)


@login_required
def master_import_approve(request, batch_id):
    if request.method != "POST":
        return HttpResponseNotAllowed(["POST"])
    batch = get_object_or_404(MasterImportBatch, pk=batch_id)
    try:
        _, counts = approve_master_import(batch.id, request.user)
    except ValidationError as exc:
        messages.error(request, " ".join(exc.messages))
    else:
        messages.success(
            request,
            "Master Product berhasil di-commit: "
            f"{counts['created']} baru, {counts['updated']} berubah, "
            f"{counts['unchanged']} tidak berubah.",
        )
    return redirect("imports:master_detail", batch_id=batch.id)


@login_required
def master_import_cancel(request, batch_id):
    if request.method != "POST":
        return HttpResponseNotAllowed(["POST"])
    batch = get_object_or_404(MasterImportBatch, pk=batch_id)
    try:
        cancel_master_import(batch.id, request.user)
    except ValidationError as exc:
        messages.error(request, " ".join(exc.messages))
    else:
        messages.success(request, "Import dibatalkan. Master Data tidak berubah.")
    return redirect("imports:master_detail", batch_id=batch.id)


@login_required
def sales_import_list(request):
    if request.method != "GET":
        return HttpResponseNotAllowed(["GET"])
    batches = SalesImportBatch.objects.exclude(
        status=SalesImportBatch.Status.VOIDED
    ).select_related("raw_file", "raw_file__uploaded_by")[:30]
    requirements = import_requirements()
    store_traffic_entry_date = timezone.localdate()
    store_traffic_entry = {
        metric.source: metric.visitors
        for metric in StoreTrafficMetric.objects.filter(
            traffic_date=store_traffic_entry_date
        )
    }
    return render(
        request,
        "imports/sales/list.html",
        {
            "batches": batches,
            "requirements": requirements,
            "requirement_summary": summarize_import_requirements(requirements),
            "can_edit_store_traffic": _can_edit_store_traffic(request.user),
            "store_traffic_entry_date": store_traffic_entry_date,
            "store_traffic_entry": store_traffic_entry,
        },
    )


def _parse_store_traffic_file(uploaded_file):
    if not uploaded_file.name.lower().endswith(".xlsx"):
        raise ValidationError("Import Traffic Toko harus memakai template .xlsx.")
    if uploaded_file.size > 5 * 1024 * 1024:
        raise ValidationError("File Traffic Toko maksimal 5 MB.")
    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except (InvalidFileException, OSError, ValueError) as exc:
        raise ValidationError("File Excel Traffic Toko tidak dapat dibaca.") from exc
    try:
        rows = workbook.active.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration as exc:
            raise ValidationError("File Traffic Toko kosong.") from exc
        columns = {
            str(value or "").strip().casefold(): index
            for index, value in enumerate(header)
        }
        required = {"tanggal", "shopee", "tiktok"}
        if not required.issubset(columns):
            raise ValidationError("Kolom template wajib: Tanggal, Shopee, TikTok.")

        result = []
        seen_dates = set()
        for row_number, row in enumerate(rows, start=2):
            if not any(value not in (None, "") for value in row):
                continue
            raw_date = row[columns["tanggal"]]
            if isinstance(raw_date, datetime):
                traffic_date = raw_date.date()
            elif isinstance(raw_date, date):
                traffic_date = raw_date
            else:
                traffic_date = None
                for date_format in ("%Y-%m-%d", "%d/%m/%Y"):
                    try:
                        traffic_date = datetime.strptime(
                            str(raw_date or "").strip(), date_format
                        ).date()
                        break
                    except ValueError:
                        continue
            if traffic_date is None:
                raise ValidationError(f"Baris {row_number}: tanggal tidak valid.")
            if traffic_date in seen_dates:
                raise ValidationError(f"Baris {row_number}: tanggal duplikat dalam file.")
            seen_dates.add(traffic_date)

            values = {}
            for source, column in (("Shopee", "shopee"), ("Tiktok", "tiktok")):
                raw_value = row[columns[column]]
                try:
                    number = Decimal(str(raw_value))
                except (InvalidOperation, TypeError, ValueError):
                    raise ValidationError(
                        f"Baris {row_number}: traffic {source} wajib angka bulat."
                    )
                if (
                    not number.is_finite()
                    or number < 0
                    or number != number.to_integral_value()
                    or number > 9223372036854775807
                ):
                    raise ValidationError(
                        f"Baris {row_number}: traffic {source} wajib angka bulat non-negatif."
                    )
                values[source] = int(number)
            result.append((traffic_date, values))
    finally:
        workbook.close()
    if not result:
        raise ValidationError("Template belum berisi data Traffic Toko.")
    return result


@login_required
def sales_store_traffic(request):
    if request.method != "POST":
        return HttpResponseNotAllowed(["POST"])
    if not _can_edit_store_traffic(request.user):
        raise PermissionDenied("Input Traffic Toko memerlukan akses Edit Sales.")
    try:
        if request.FILES.get("traffic_file"):
            rows = _parse_store_traffic_file(request.FILES["traffic_file"])
            with transaction.atomic():
                for traffic_date, values in rows:
                    _save_store_traffic(
                        request, traffic_date=traffic_date, values=values
                    )
        else:
            rows = None
            _save_store_traffic(request)
    except ValidationError as exc:
        messages.error(request, exc.messages[0])
    else:
        message = (
            f"Traffic toko berhasil di-import untuk {len(rows)} tanggal."
            if rows is not None
            else "Traffic toko harian berhasil disimpan."
        )
        messages.success(request, message)
    return redirect("imports:sales_list")


@login_required
def sales_store_traffic_template(request):
    if not _can_edit_store_traffic(request.user):
        raise PermissionDenied("Template Traffic Toko memerlukan akses Edit Sales.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Traffic Toko"
    sheet.append(["Tanggal", "Shopee", "TikTok"])
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 16
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        'attachment; filename="template-traffic-toko.xlsx"'
    )
    return response


@login_required
def sales_import_upload(request):
    if request.method == "POST":
        form = SalesImportUploadForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                batch = create_sales_import(
                    form.cleaned_data["file"],
                    form.cleaned_data["source"],
                    request.user,
                )
            except DuplicateRawFile as exc:
                existing_batch = exc.raw_file.sales_batches.order_by("-created_at").first()
                detail = f" Batch sebelumnya: {existing_batch.id}." if existing_batch else ""
                form.add_error("file", "File identik sudah pernah diunggah." + detail)
            else:
                if batch.status == SalesImportBatch.Status.READY:
                    messages.success(request, "File Sales berhasil diparsing dan siap direview.")
                else:
                    messages.warning(request, "File diparsing tetapi memiliki blocking issue.")
                return redirect("imports:sales_detail", batch_id=batch.id)
    elif request.method == "GET":
        form = SalesImportUploadForm()
    else:
        return HttpResponseNotAllowed(["GET", "POST"])
    return render(request, "imports/sales/upload.html", {"form": form})


@login_required
def sales_import_detail(request, batch_id):
    batch = get_object_or_404(
        SalesImportBatch.objects.select_related("raw_file", "approved_by"),
        pk=batch_id,
    )
    action_filter = request.GET.get("action", "")
    severity_filter = request.GET.get("severity", "")
    staged_rows = batch.staged_rows.select_related("sku", "existing_line")
    if action_filter in {choice for choice, _ in StagedSalesRow.ProposedAction.choices}:
        staged_rows = staged_rows.filter(proposed_action=action_filter)
    page = Paginator(staged_rows, 50).get_page(request.GET.get("page"))

    issues = batch.issues.select_related("staged_row")
    if severity_filter in {choice for choice, _ in SalesImportIssue.Severity.choices}:
        issues = issues.filter(severity=severity_filter)
    return render(
        request,
        "imports/sales/detail.html",
        {
            "batch": batch,
            "page": page,
            "issues": issues[:100],
            "commit_enabled": settings.SALES_IMPORT_COMMIT_ENABLED,
            "fifo_opening_ready": FIFOOpeningImportBatch.objects.filter(status=FIFOOpeningImportBatch.Status.COMMITTED).exists(),
        },
    )


@login_required
def sales_import_approve(request, batch_id):
    if request.method != "POST":
        return HttpResponseNotAllowed(["POST"])
    batch = get_object_or_404(SalesImportBatch, pk=batch_id)
    try:
        _, counts = approve_sales_import(batch.id, request.user)
    except ValidationError as exc:
        messages.error(request, " ".join(exc.messages))
    else:
        messages.success(
            request,
            f"Sales committed: {counts['orders_created']} order baru, "
            f"{counts['lines_created']} line baru, {counts['status_updates']} status berubah.",
        )
    return redirect("imports:sales_detail", batch_id=batch.id)


@login_required
def sales_import_reparse(request, batch_id):
    if request.method != "POST":
        return HttpResponseNotAllowed(["POST"])
    batch = get_object_or_404(SalesImportBatch, pk=batch_id)
    if batch.status not in {SalesImportBatch.Status.BLOCKED, SalesImportBatch.Status.READY}:
        messages.error(request, "Hanya batch yang belum di-commit yang dapat dicek ulang.")
        return redirect("imports:sales_detail", batch_id=batch.id)
    batch = parse_sales_batch(batch)
    if batch.status == SalesImportBatch.Status.READY:
        messages.success(request, "File lama berhasil dicek ulang dan sekarang siap direview.")
    else:
        messages.warning(request, "File sudah dicek ulang, tetapi masih memiliki blocking issue lain.")
    return redirect("imports:sales_detail", batch_id=batch.id)
