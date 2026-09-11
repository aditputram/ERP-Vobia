from decimal import Decimal
from io import BytesIO
from urllib.parse import urlencode

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.db import transaction
from django.http import Http404, HttpResponse, HttpResponseNotAllowed
from django.db.models import Q, Sum
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.dateparse import parse_date
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from master_data.models import Category, Product, ProductStatus, SKU, Warehouse
from purchasing.models import PurchaseOrder, PurchaseOrderLine
from sales.models import SalesOrder

from imports.services.storage import DuplicateRawFile

from .forms import AdjustmentForm, DeliveryReceiveForm, FIFOOpeningImportUploadForm, InboundForm, OpeningForm, QCForm, ReturnForm, WarehouseForm
from .documents import (
    build_inbound_checklist_pdf,
    build_inbound_filtered_receipts_pdf,
    build_inbound_queue_pdf,
    build_inbound_receipt_pdf,
)
from .models import (
    FIFOLayer,
    FIFOOpeningImportBatch,
    FIFOOpeningImportIssue,
    InboundReceipt,
    InventoryException,
    InventoryMovement,
    PhysicalReturnReceipt,
    QCInspection,
)
from .services.aging import po_aging_snapshot, refresh_po_close
from .services.fifo import CUTOVER_DATE, create_expected_return, inventory_balance, post_adjustment, post_opening, qc_approved_qty, receive_rejected_goods_delivery, record_inbound, record_physical_return, record_qc
from .services.opening_import import approve_opening_import, create_opening_import
from .services.reporting import filtered_skus, inventory_parent_summary_rows, inventory_summary_rows, movement_ledger_rows, parent_movement_ledger_rows
from production.models import ProductionActivity
from config.excel_exports import append_table, workbook_response


def _excel_text(value):
    value = str(value or "")
    return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value


def _export_inventory(balances, *, as_of_date, warehouse, sku_type, stock_status):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inventory Summary"
    common = (
        "As Of Date", "Warehouse", "Ending 31 Juli", "Movement In", "Movement Out",
        "Ending Stock", "FIFO Remaining", "FIFO Value", "Active Exceptions", "Stock Status",
        "Warehouse Actual Qty", "Evidence Reference", "Warehouse Notes",
    )
    identity = (
        ("Parent SKU", "Product", "Category", "SKU Count")
        if sku_type == "parent"
        else ("SKU", "Parent SKU", "Product", "Variant", "Category")
    )
    headers = common[:2] + identity + common[2:]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"

    for row in balances:
        product = row["product"] if sku_type == "parent" else row["sku"].product_variant.product
        identity_values = (
            (
                _excel_text(row["parent_sku"]), _excel_text(product.name),
                _excel_text(product.category.name), row["sku_count"],
            )
            if sku_type == "parent"
            else (
                _excel_text(row["sku"].sku), _excel_text(product.parent_sku),
                _excel_text(product.name), _excel_text(row["sku"].product_variant.name),
                _excel_text(product.category.name),
            )
        )
        sheet.append((
            as_of_date, _excel_text(warehouse.name if warehouse else "All Warehouse"),
            *identity_values,
            row["opening_qty"], row["incoming_qty"], row["outgoing_qty"], row["balance"],
            row["fifo_qty"], row["fifo_value"], row["exception_count"], row["stock_status"],
            "", "", "",
        ))

    sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate((14, 20, 20, 20, 30, 22, 20, 16, 16, 16, 16, 18, 18, 18, 16, 22, 28, 28), start=1):
        if index <= sheet.max_column:
            sheet.column_dimensions[get_column_letter(index)].width = width
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    status_label = (stock_status or "all").lower()
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="VOBIA-Inventory-{status_label}-{as_of_date:%Y-%m-%d}.xlsx"'
    )
    return response


def _export_turnover(rows, *, sku_type):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inventory Turnover"
    append_table(
        sheet,
        (
            "Date", "Warehouse", "Parent SKU" if sku_type == "parent" else "SKU", "Movement",
            "Direction", "Qty", "Signed Qty", "Allocated Cost", "Balance Qty", "Balance Value",
            "Reference", "Ledger Key",
        ),
        (
            (
                row["date"],
                row["warehouse_name"],
                row["parent_sku"] if sku_type == "parent" else row["sku"].sku,
                row["type_label"],
                row["direction"],
                row["quantity"],
                row["signed_quantity"],
                row["allocated_cost"],
                row["running_balance"],
                row["running_value"],
                row["reference"],
                row["key"],
            )
            for row in rows
        ),
        number_formats={6: "#,##0", 7: "#,##0", 8: "Rp #,##0", 9: "#,##0", 10: "Rp #,##0"},
    )
    return workbook_response(workbook, f"VOBIA-Inventory-Turnover-{timezone.localdate():%Y-%m-%d}.xlsx")


def _export_inbound(delivery_orders, completed_delivery_orders, outstanding, receipts):
    workbook = Workbook()
    queue_sheet = workbook.active
    queue_sheet.title = "Delivery Queue"
    append_table(
        queue_sheet,
        ("No. DO", "Jenis", "Tanggal Kirim", "PO", "Vendor", "Jumlah SKU", "Qty Delivering", "Qty Received", "Remaining"),
        (
            (
                shipment["delivery_order"].number,
                shipment["kind_label"],
                shipment["delivery_date"],
                shipment["production_order"].po.po_number,
                shipment["production_order"].po.supplier.name,
                len(shipment["rows"]),
                shipment["remaining"],
                shipment["received"],
                shipment["remaining"],
            )
            for shipment in delivery_orders
        ),
        number_formats={6: "#,##0", 7: "#,##0", 8: "#,##0", 9: "#,##0"},
    )
    completed_sheet = workbook.create_sheet("Completed Delivery")
    append_table(
        completed_sheet,
        ("No. DO", "Jenis", "Tanggal Kirim", "Tanggal Diterima", "PO", "Vendor", "Qty Dikirim", "Qty Received", "Status"),
        (
            (
                shipment["delivery_order"].number,
                shipment["kind_label"],
                shipment["delivery_date"],
                shipment["received_date"],
                shipment["production_order"].po.po_number,
                shipment["production_order"].po.supplier.name,
                shipment["shipped"],
                shipment["received"],
                "Received",
            )
            for shipment in completed_delivery_orders
        ),
        number_formats={7: "#,##0", 8: "#,##0"},
    )
    outstanding_sheet = workbook.create_sheet("Outstanding PO")
    append_table(
        outstanding_sheet,
        ("PO", "Source", "SKU", "Product", "Variant", "Size", "Ordered", "QC Passed", "Received", "Eligible Now", "Outstanding"),
        (
            (
                row["line"].po.po_number,
                row["line"].po.get_source_display(),
                row["line"].sku.sku,
                row["line"].sku.product_variant.product.name,
                row["line"].sku.product_variant.name,
                row["line"].sku.size,
                row["line"].ordered_qty,
                row["qc_passed"],
                row["received"],
                row["eligible"],
                row["outstanding"],
            )
            for row in outstanding
        ),
        number_formats={column: "#,##0" for column in range(7, 12)},
    )
    history_sheet = workbook.create_sheet("Inbound History")
    append_table(
        history_sheet,
        ("Date", "No. DO", "Reference", "PO", "SKU", "Product", "Variant", "Size", "Qty", "COGS Snapshot", "Retail Snapshot", "Warehouse", "Recorded By"),
        (
            (
                row.inbound_date,
                (
                    row.delivery_activity.delivery_order.number
                    if row.delivery_activity_id and row.delivery_activity.delivery_order_id
                    else ""
                ),
                row.reference,
                row.po_line.po.po_number,
                row.po_line.sku.sku,
                row.po_line.sku.product_variant.product.name,
                row.po_line.sku.product_variant.name,
                row.po_line.sku.size,
                row.received_qty,
                row.po_line.cogs_snapshot,
                row.retail_price_snapshot,
                row.warehouse.name,
                row.recorded_by.username,
            )
            for row in receipts
        ),
        number_formats={9: "#,##0", 10: "Rp #,##0", 11: "Rp #,##0"},
    )
    return workbook_response(workbook, f"VOBIA-Inbound-{timezone.localdate():%Y-%m-%d}.xlsx")


def _export_returns(return_rows, claim_receipts, receipts, selected_order):
    workbook = Workbook()
    selected_sheet = workbook.active
    selected_sheet.title = "Selected Order"
    append_table(
        selected_sheet,
        ("Source", "Order", "SKU", "Product", "Variant", "Size", "Sales Qty", "Expected Return", "Received", "Remaining"),
        (
            (
                selected_order.display_source if selected_order else "",
                selected_order.order_number if selected_order else "",
                row["line"].sku.sku,
                row["line"].sku.product_variant.product.name,
                row["line"].sku.product_variant.name,
                row["line"].sku.size,
                row["line"].quantity,
                row["expected_qty"],
                row["returned_qty"],
                row["remaining_qty"],
            )
            for row in return_rows
        ),
        number_formats={column: "#,##0" for column in range(7, 11)},
    )
    claims_sheet = workbook.create_sheet("Marketplace Claims")
    append_table(
        claims_sheet,
        ("Date", "Source", "Order", "SKU", "Product", "Qty", "Condition", "Follow-up Status", "Warehouse", "Received By"),
        (
            (
                row.received_date,
                row.sales_line.order.display_source,
                row.sales_line.order.order_number,
                row.sales_line.sku.sku,
                row.sales_line.sku.product_variant.product.name,
                row.quantity,
                row.get_condition_display(),
                row.get_follow_up_status_display(),
                row.warehouse.name,
                row.recorded_by.username,
            )
            for row in claim_receipts
        ),
        number_formats={6: "#,##0"},
    )
    history_sheet = workbook.create_sheet("Return History")
    append_table(
        history_sheet,
        ("Date", "Source", "Order", "SKU", "Product", "Qty", "Condition", "Stock Impact", "Warehouse", "Received By"),
        (
            (
                row.received_date,
                row.sales_line.order.display_source,
                row.sales_line.order.order_number,
                row.sales_line.sku.sku,
                row.sales_line.sku.product_variant.product.name,
                row.quantity,
                row.get_condition_display(),
                "FIFO restored" if row.condition == PhysicalReturnReceipt.Condition.SELLABLE else "No stock",
                row.warehouse.name,
                row.recorded_by.username,
            )
            for row in receipts
        ),
        number_formats={6: "#,##0"},
    )
    return workbook_response(workbook, f"VOBIA-Return-Log-{timezone.localdate():%Y-%m-%d}.xlsx")


def _export_outbound(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Outbound"
    append_table(
        sheet,
        ("Date", "Source", "Order", "SKU", "Qty", "FIFO COGS", "Status", "Ledger Key", "Posted By"),
        (
            (
                row.movement_date,
                row.sales_line.order.display_source,
                row.sales_line.order.order_number,
                row.sku.sku,
                row.quantity,
                row.allocated_cost,
                row.sales_line.order.current_status,
                row.movement_key,
                row.posted_by.username,
            )
            for row in rows
        ),
        number_formats={5: "#,##0", 6: "Rp #,##0"},
    )
    return workbook_response(workbook, f"VOBIA-Outbound-{timezone.localdate():%Y-%m-%d}.xlsx")


@login_required
def production(request):
    return redirect("production:dashboard")


@login_required
def overview(request):
    forms = {"warehouse": WarehouseForm(), "adjustment": AdjustmentForm()}
    if request.method == "POST":
        form_name = request.POST.get("form_name")
        if form_name == "warehouse":
            forms["warehouse"] = WarehouseForm(request.POST)
            if forms["warehouse"].is_valid():
                forms["warehouse"].save()
                messages.success(request, "Warehouse berhasil ditambahkan.")
                return redirect("inventory:overview")
        elif form_name == "adjustment":
            forms["adjustment"] = AdjustmentForm(request.POST)
            if forms["adjustment"].is_valid():
                data = forms["adjustment"].cleaned_data
                try:
                    post_adjustment(actor=request.user, **data)
                except ValidationError as exc:
                    forms["adjustment"].add_error(None, exc)
                else:
                    messages.success(request, "Adjustment traceable berhasil diposting dan exception terkait diperbarui.")
                    return redirect("inventory:overview")
        elif form_name == "refresh_aging":
            for po in PurchaseOrder.objects.filter(status=PurchaseOrder.Status.RELEASED):
                refresh_po_close(po.id)
            messages.success(request, "PO Aging dan close/reopen condition diperbarui.")
            return redirect("inventory:overview")
    query = request.GET.get("q", "").strip()
    status = request.GET.get("status", "")
    category = request.GET.get("category", "")
    stock_status = request.GET.get("stock_status", "")
    warehouses = list(Warehouse.objects.filter(is_active=True, code__in=("MAIN", "REJECT")).order_by("name"))
    warehouse_map = {str(row.id): row for row in warehouses}
    main_warehouse = next((row for row in warehouses if row.code == "MAIN"), None)
    warehouse_value = request.GET.get("warehouse")
    selected_warehouse = (
        None
        if "warehouse" in request.GET and warehouse_value == ""
        else warehouse_map.get(warehouse_value, main_warehouse)
    )
    sku_type = request.GET.get("sku_type", "sku")
    if sku_type not in {"sku", "parent"}:
        sku_type = "sku"
    requested_as_of_date = request.GET.get("as_of_date", "").strip()
    today = timezone.localdate()
    date_filter_error = ""
    if requested_as_of_date:
        as_of_date = parse_date(requested_as_of_date)
        if as_of_date is None:
            as_of_date = today
            date_filter_error = "Tanggal tidak valid; posisi stok dikembalikan ke hari ini."
    else:
        as_of_date = today
    if as_of_date < CUTOVER_DATE:
        as_of_date = CUTOVER_DATE
        date_filter_error = "Riwayat Warehouse ERP tersedia mulai FIFO cutover 31 July 2026."
    skus = filtered_skus(query=query, status=status, category=category)
    balances = inventory_summary_rows(skus, as_of_date=as_of_date, warehouse=selected_warehouse)
    if selected_warehouse:
        balances = [
            row
            for row in balances
            if any(row[field] for field in ("opening_qty", "incoming_qty", "outgoing_qty", "fifo_qty"))
        ]
    if sku_type == "parent":
        balances = inventory_parent_summary_rows(balances)
    if stock_status:
        balances = [row for row in balances if row["stock_status"] == stock_status]
    if request.method == "GET" and request.GET.get("export") == "xlsx":
        return _export_inventory(
            balances,
            as_of_date=as_of_date,
            warehouse=selected_warehouse,
            sku_type=sku_type,
            stock_status=stock_status,
        )
    total_balance = sum((row["balance"] for row in balances), 0)
    total_fifo_value = sum((row["fifo_value"] for row in balances), 0)
    total_exceptions = sum((row["exception_count"] for row in balances), 0)
    pos = list(PurchaseOrder.objects.filter(status=PurchaseOrder.Status.RELEASED).prefetch_related("lines")[:100])
    for po in pos:
        po.aging_snapshot = po_aging_snapshot(po)
    return render(
        request,
        "inventory/overview.html",
        {
            **forms,
            "balances": balances,
            "exceptions": InventoryException.objects.filter(status=InventoryException.Status.OPEN).select_related("sku", "movement")[:200],
            "pos": pos,
            "opening_batch": FIFOOpeningImportBatch.objects.first(),
            "query": query,
            "selected_status": status,
            "selected_category": category,
            "selected_stock_status": stock_status,
            "warehouses": warehouses,
            "selected_warehouse": selected_warehouse,
            "sku_type": sku_type,
            "as_of_date": as_of_date,
            "cutover_date": CUTOVER_DATE,
            "date_filter_error": date_filter_error,
            "product_statuses": ProductStatus.objects.filter(is_active=True),
            "categories": Category.objects.filter(is_active=True),
            "total_balance": total_balance,
            "total_fifo_value": total_fifo_value,
            "total_exceptions": total_exceptions,
            "negative_sku_count": sum(1 for row in balances if row["balance"] < 0),
        },
    )


@login_required
def turnover(request):
    query = request.GET.get("q", "").strip()
    sku_type = request.GET.get("sku_type", "sku")
    if sku_type not in {"sku", "parent"}:
        sku_type = "sku"
    products = list(Product.objects.filter(is_active=True, variants__skus__is_active=True).distinct())
    product_options = [{"value": str(row.id), "label": row.name} for row in products]
    valid_product_ids = {row["value"] for row in product_options}
    selected_products = [value for value in request.GET.getlist("product") if value in valid_product_ids]
    size_queryset = SKU.objects.filter(is_active=True).exclude(size="")
    if selected_products:
        size_queryset = size_queryset.filter(product_variant__product_id__in=selected_products)
    size_values = list(size_queryset.values_list("size", flat=True).distinct().order_by("size"))
    selected_sizes = [value for value in request.GET.getlist("size") if value in size_values]
    movement_type = request.GET.get("movement_type", "")
    warehouses = list(Warehouse.objects.filter(is_active=True, code__in=("MAIN", "REJECT")).order_by("name"))
    warehouse_map = {str(row.id): row for row in warehouses}
    main_warehouse = next((row for row in warehouses if row.code == "MAIN"), None)
    warehouse_value = request.GET.get("warehouse")
    selected_warehouse = (
        None
        if "warehouse" in request.GET and warehouse_value == ""
        else warehouse_map.get(warehouse_value, main_warehouse)
    )
    date_from_value = request.GET.get("date_from", "")
    date_to_value = request.GET.get("date_to", "")
    date_from = parse_date(date_from_value) if date_from_value else None
    date_to = parse_date(date_to_value) if date_to_value else None
    skus = filtered_skus(query=query, product=selected_products, size=selected_sizes)
    valid_movement_type = movement_type if movement_type in InventoryMovement.MovementType.values else ""
    if sku_type == "parent":
        rows = parent_movement_ledger_rows(
            movement_ledger_rows(skus, date_to=date_to, warehouse=selected_warehouse)
        )
        if date_from:
            rows = [row for row in rows if row["date"] >= date_from]
        if valid_movement_type:
            rows = [row for row in rows if row["type"] == valid_movement_type]
    else:
        rows = movement_ledger_rows(
            skus,
            date_from=date_from,
            date_to=date_to,
            movement_type=valid_movement_type,
            warehouse=selected_warehouse,
        )
    if request.GET.get("export") == "xlsx":
        return _export_turnover(rows, sku_type=sku_type)
    page = Paginator(rows, 100).get_page(request.GET.get("page"))
    return render(request, "inventory/turnover.html", {
        "page": page,
        "query": query,
        "sku_type": sku_type,
        "selected_products": selected_products,
        "selected_sizes": selected_sizes,
        "product_options": product_options,
        "size_options": [{"value": value, "label": value} for value in size_values],
        "movement_type": movement_type,
        "date_from": date_from_value,
        "date_to": date_to_value,
        "movement_types": InventoryMovement.MovementType.choices,
        "warehouses": warehouses,
        "selected_warehouse": selected_warehouse,
        "total_rows": len(rows),
    })


@login_required
def inbound(request):
    query = request.GET.get("q", "").strip()
    requested_inbound_date = parse_date(
        request.POST.get("inbound_date", "")
        if request.method == "POST" and request.POST.get("form_name") == "delivery_receive"
        else request.GET.get("inbound_date", "")
    )
    delivery_rows = []
    deliveries = ProductionActivity.objects.filter(
        entry_kind=ProductionActivity.EntryKind.ACTIVITY,
        activity_type__in=(
            ProductionActivity.ActivityType.WAREHOUSE_DELIVERY,
            ProductionActivity.ActivityType.REJECTED_WAREHOUSE_DELIVERY,
        ),
    ).select_related(
        "production_order__po__supplier",
        "po_line__sku__product_variant__product",
        "actor",
        "delivery_order",
    ).prefetch_related(
        "correction_entries",
        "inbound_receipts",
        "rejected_follow_ups__received_warehouse",
    ).order_by("-activity_date", "-occurred_at")
    for delivery in deliveries:
        effective = max(delivery.correction_entries.all(), key=lambda row: row.occurred_at, default=delivery)
        shipped = effective.quantity or 0
        is_rejected = delivery.activity_type == ProductionActivity.ActivityType.REJECTED_WAREHOUSE_DELIVERY
        rejected_follow_ups = list(delivery.rejected_follow_ups.all())
        received = (
            sum(
                (row.open_qty for row in rejected_follow_ups if row.delivery_status == "INBOUND"),
                Decimal("0"),
            )
            if is_rejected
            else sum((receipt.received_qty for receipt in delivery.inbound_receipts.all()), 0)
        )
        remaining = max(shipped - received, 0)
        latest_receipt = (
            None
            if is_rejected
            else max(delivery.inbound_receipts.all(), key=lambda receipt: receipt.created_at, default=None)
        )
        rejected_receipt = next((row for row in rejected_follow_ups if row.received_date), None)
        delivery_rows.append(
            {
                "delivery": delivery,
                "is_rejected": is_rejected,
                "kind_label": "Rejected Goods" if is_rejected else "QC Passed",
                "shipped": shipped,
                "received": received,
                "remaining": remaining,
                "latest_receipt": latest_receipt,
                "received_date": rejected_receipt.received_date if rejected_receipt else getattr(latest_receipt, "inbound_date", None),
                "received_warehouse": rejected_receipt.received_warehouse if rejected_receipt else getattr(latest_receipt, "warehouse", None),
            }
        )

    selected_delivery_id = request.POST.get("delivery_activity", "")
    selected_delivery = next(
        (
            row
            for row in delivery_rows
            if row["remaining"] and str(row["delivery"].id) == selected_delivery_id
        ),
        None,
    )
    is_delivery_receive = request.method == "POST" and request.POST.get("form_name") == "delivery_receive"
    delivery_form = None
    if selected_delivery:
        delivery_form = DeliveryReceiveForm(
            request.POST if is_delivery_receive else None,
            max_qty=selected_delivery["remaining"],
            min_date=selected_delivery["delivery"].activity_date,
            initial={"delivery_activity": selected_delivery["delivery"].id},
            default_warehouse_code="REJECT" if selected_delivery["is_rejected"] else "MAIN",
        )
    if is_delivery_receive and selected_delivery and delivery_form.is_valid():
        values = delivery_form.cleaned_data.copy()
        values.pop("delivery_activity")
        try:
            if selected_delivery["is_rejected"]:
                receive_rejected_goods_delivery(
                    delivery_activity=selected_delivery["delivery"],
                    actor=request.user,
                    **values,
                )
            else:
                values["reference"] = (
                    f"{selected_delivery['delivery'].delivery_order.number}/"
                    f"{selected_delivery['delivery'].id}/"
                    f"{selected_delivery['delivery'].inbound_receipts.count() + 1:03d}"
                )
                record_inbound(
                    po_line=selected_delivery["delivery"].po_line,
                    delivery_activity=selected_delivery["delivery"],
                    actor=request.user,
                    **values,
                )
        except ValidationError as exc:
            delivery_form.add_error(None, exc)
        else:
            messages.success(
                request,
                "Rejected Goods diterima tanpa menambah stock/FIFO."
                if selected_delivery["is_rejected"]
                else "Pengiriman diterima; Delivering berkurang dan Inbound bertambah.",
            )
            delivery_order_id = selected_delivery["delivery"].delivery_order_id
            return redirect(
                f"{reverse('inventory:inbound')}?delivery_order={delivery_order_id}"
                f"&inbound_date={values['inbound_date']:%Y-%m-%d}"
                f"#delivery-order-{delivery_order_id}"
            )

    delivery_order_map = {}
    open_delivery_order_id = request.GET.get("delivery_order", "")
    for row in delivery_rows:
        delivery = row["delivery"]
        row["receive_form"] = None
        if row["remaining"]:
            row["receive_form"] = (
                delivery_form
                if selected_delivery and delivery.id == selected_delivery["delivery"].id
                else DeliveryReceiveForm(
                    max_qty=row["remaining"],
                    min_date=delivery.activity_date,
                    initial={
                        "delivery_activity": delivery.id,
                        "inbound_date": (
                            requested_inbound_date
                            if requested_inbound_date and requested_inbound_date >= delivery.activity_date
                            else max(timezone.localdate(), delivery.activity_date)
                        ),
                    },
                    default_warehouse_code="REJECT" if row["is_rejected"] else "MAIN",
                )
            )
            row["receive_form_id"] = f"delivery-receive-{delivery.id}"
            for field in row["receive_form"].fields.values():
                field.widget.attrs["form"] = row["receive_form_id"]
        key = str(delivery.delivery_order_id)
        group = delivery_order_map.setdefault(
            key,
            {
                "delivery_order": delivery.delivery_order,
                "production_order": delivery.production_order,
                "delivery_date": delivery.activity_date,
                "received_date": None,
                "actor": delivery.actor,
                "kind_label": row["kind_label"],
                "rows": [],
                "shipped": 0,
                "received": 0,
                "remaining": 0,
                "open": key == open_delivery_order_id,
            },
        )
        group["rows"].append(row)
        group["shipped"] += row["shipped"]
        group["received"] += row["received"]
        group["remaining"] += row["remaining"]
        if row["received_date"] and (
            group["received_date"] is None
            or row["received_date"] > group["received_date"]
        ):
            group["received_date"] = row["received_date"]
        if selected_delivery and delivery.id == selected_delivery["delivery"].id:
            group["open"] = True
    delivery_orders = [group for group in delivery_order_map.values() if group["remaining"]]
    completed_delivery_orders = [
        group for group in delivery_order_map.values() if not group["remaining"] and group["received"]
    ]
    if query:
        terms = query.casefold().split()

        def matches(group):
            haystack = " ".join(
                [group["production_order"].po.po_number]
                + [
                    f'{row["delivery"].po_line.sku.sku} {row["delivery"].po_line.sku.product_variant.product.name}'
                    for row in group["rows"]
                ]
            )
            return all(term in haystack.casefold() for term in terms)

        completed_delivery_orders = [group for group in completed_delivery_orders if matches(group)]

    form = InboundForm(request.POST if request.method == "POST" and not is_delivery_receive else None)
    if request.method == "POST" and not is_delivery_receive and form.is_valid():
        try:
            record_inbound(actor=request.user, **form.cleaned_data)
        except ValidationError as exc:
            form.add_error(None, exc)
        else:
            messages.success(request, "Inbound aktual diposting; movement dan FIFO cost layer terbentuk.")
            return redirect("inventory:inbound")
    outstanding = []
    lines = PurchaseOrderLine.objects.filter(po__status=PurchaseOrder.Status.RELEASED).select_related(
        "po", "po__supplier", "sku"
    ).prefetch_related("qc_inspections", "inbound_receipts")
    for line in lines:
        qc_passed = qc_approved_qty(line)
        received = line.received_before_cutover_qty + sum((row.received_qty for row in line.inbound_receipts.all()), 0)
        outstanding_qty = max(line.ordered_qty - received, 0)
        eligible_qty = max(qc_passed - received, 0)
        if outstanding_qty:
            outstanding.append({"line": line, "qc_passed": qc_passed, "received": received, "outstanding": outstanding_qty, "eligible": eligible_qty})
    receipt_rows = InboundReceipt.objects.select_related(
        "po_line__po",
        "po_line__sku__product_variant__product",
        "warehouse",
        "recorded_by",
        "delivery_activity__delivery_order",
    ).order_by("-inbound_date", "-created_at")
    if request.method == "GET" and request.GET.get("export") == "xlsx":
        return _export_inbound(delivery_orders, completed_delivery_orders, outstanding, receipt_rows)
    receipts = receipt_rows[:300]
    return render(
        request,
        "inventory/inbound.html",
        {
            "form": form,
            "delivery_rows": delivery_rows,
            "delivery_orders": delivery_orders,
            "completed_delivery_orders": completed_delivery_orders,
            "selected_delivery": selected_delivery,
            "delivery_form": delivery_form,
            "outstanding": outstanding,
            "receipts": receipts,
            "query": query,
        },
    )


@login_required
def inbound_po_receipt_pdf(request, po_id):
    po = get_object_or_404(
        PurchaseOrder.objects.select_related("supplier").prefetch_related("lines"),
        pk=po_id,
    )
    receipts = InboundReceipt.objects.filter(po_line__po=po).select_related(
        "po_line__sku__product_variant__product",
        "warehouse",
        "recorded_by",
        "delivery_activity__delivery_order",
    ).order_by("inbound_date", "created_at")
    if not receipts.exists():
        raise Http404("PO belum memiliki actual inbound.")
    response = HttpResponse(
        build_inbound_receipt_pdf(po=po, receipts=receipts),
        content_type="application/pdf",
    )
    safe_number = po.po_number.replace("/", "-")
    response["Content-Disposition"] = f'attachment; filename="VOBIA-Inbound-{safe_number}.pdf"'
    return response


@login_required
def inbound_po_checklist_pdf(request, po_id):
    po = get_object_or_404(PurchaseOrder.objects.select_related("supplier"), pk=po_id)
    activities = ProductionActivity.objects.filter(
        po_line__po=po,
        entry_kind=ProductionActivity.EntryKind.ACTIVITY,
        activity_type=ProductionActivity.ActivityType.WAREHOUSE_DELIVERY,
    ).select_related("po_line__sku__product_variant__product", "delivery_order").prefetch_related("correction_entries", "inbound_receipts").order_by("activity_date", "occurred_at")
    deliveries = []
    for delivery in activities:
        effective = max(delivery.correction_entries.all(), key=lambda row: row.occurred_at, default=delivery)
        shipped = effective.quantity or 0
        received = sum((receipt.received_qty for receipt in delivery.inbound_receipts.all()), 0)
        if shipped > received:
            deliveries.append({"delivery": delivery, "shipped": shipped, "received": received, "remaining": shipped - received})
    if not deliveries:
        raise Http404("PO tidak memiliki pengiriman yang menunggu penerimaan.")
    response = HttpResponse(build_inbound_checklist_pdf(po=po, deliveries=deliveries), content_type="application/pdf")
    safe_number = po.po_number.replace("/", "-")
    response["Content-Disposition"] = f'attachment; filename="VOBIA-Checklist-Inbound-{safe_number}.pdf"'
    return response


@login_required
def inbound_pending_report_pdf(request):
    activities = ProductionActivity.objects.filter(
        entry_kind=ProductionActivity.EntryKind.ACTIVITY,
        activity_type__in=(ProductionActivity.ActivityType.WAREHOUSE_DELIVERY, ProductionActivity.ActivityType.REJECTED_WAREHOUSE_DELIVERY),
    ).select_related("production_order__po__supplier", "po_line__sku__product_variant__product", "delivery_order").prefetch_related("correction_entries", "inbound_receipts", "rejected_follow_ups").order_by("activity_date", "occurred_at")
    groups = {}
    for delivery in activities:
        effective = max(delivery.correction_entries.all(), key=lambda row: row.occurred_at, default=delivery)
        shipped = effective.quantity or 0
        is_rejected = delivery.activity_type == ProductionActivity.ActivityType.REJECTED_WAREHOUSE_DELIVERY
        received = sum((row.open_qty for row in delivery.rejected_follow_ups.all() if row.delivery_status == "INBOUND"), Decimal("0")) if is_rejected else sum((row.received_qty for row in delivery.inbound_receipts.all()), 0)
        remaining = max(shipped - received, 0)
        if not remaining:
            continue
        group = groups.setdefault(str(delivery.delivery_order_id), {"rows": []})
        group["rows"].append({"delivery": delivery, "kind_label": "Rejected Goods" if is_rejected else "QC Passed", "remaining": remaining})
    deliveries = list(groups.values())
    if not deliveries:
        raise Http404("Tidak ada pengiriman yang menunggu penerimaan.")
    response = HttpResponse(build_inbound_queue_pdf(deliveries=deliveries), content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="VOBIA-Pengiriman-Menunggu-Penerimaan.pdf"'
    return response


@login_required
def inbound_completed_report_pdf(request):
    query = request.GET.get("q", "").strip()
    terms = query.casefold().split()
    receipts = InboundReceipt.objects.select_related(
        "po_line__po__supplier", "po_line__sku__product_variant__product", "warehouse", "recorded_by", "delivery_activity__delivery_order"
    ).order_by("inbound_date", "created_at")
    if terms:
        receipts = [
            row for row in receipts
            if all(term in f"{row.po_line.po.po_number} {row.po_line.sku.sku} {row.po_line.sku.product_variant.product.name}".casefold() for term in terms)
        ]
    if not receipts:
        raise Http404("Tidak ada pengiriman diterima yang sesuai filter.")
    response = HttpResponse(build_inbound_filtered_receipts_pdf(receipts=receipts, query=query), content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="VOBIA-Pengiriman-Sudah-Diterima.pdf"'
    return response


@login_required
def return_log(request):
    return_orders = SalesOrder.objects.filter(
        current_status="Retur",
        lines__sku__isnull=False,
    ).distinct()
    source_options = sorted(
        {source_label or source for source_label, source in return_orders.values_list("source_label", "source")},
        key=str.casefold,
    )
    selected_source = request.GET.get("source", "").strip()
    if selected_source not in source_options:
        selected_source = ""
    filtered_orders = return_orders.none()
    if selected_source:
        filtered_orders = return_orders.filter(
            Q(source_label=selected_source) | Q(source_label="", source=selected_source)
        ).order_by("-order_datetime", "order_number")

    selected_order_number = request.GET.get("order_number", "").strip()
    selected_order = (
        filtered_orders.filter(order_number=selected_order_number).first()
        if selected_order_number
        else None
    )
    return_rows = []
    if selected_order:
        lines = selected_order.lines.filter(sku__isnull=False).select_related(
            "sku__product_variant__product",
            "expected_return",
        ).annotate(returned_qty=Sum("physical_returns__quantity"))
        for line in lines:
            expected = getattr(line, "expected_return", None)
            expected_qty = expected.expected_qty if expected else Decimal(line.quantity)
            returned_qty = line.returned_qty or Decimal("0")
            return_rows.append(
                {
                    "line": line,
                    "expected_qty": expected_qty,
                    "returned_qty": returned_qty,
                    "remaining_qty": max(expected_qty - returned_qty, Decimal("0")),
                }
            )

    form = ReturnForm(request.POST or None)
    submitted_quantities = {}
    claim_conditions = [
        PhysicalReturnReceipt.Condition.DAMAGED,
        PhysicalReturnReceipt.Condition.DEFECTIVE,
        PhysicalReturnReceipt.Condition.MISSING,
        PhysicalReturnReceipt.Condition.WRONG_ITEM,
    ]
    if request.method == "POST" and request.POST.get("action") == "update_follow_up":
        receipt = get_object_or_404(
            PhysicalReturnReceipt,
            pk=request.POST.get("receipt_id"),
            condition__in=claim_conditions,
        )
        follow_up_status = request.POST.get("follow_up_status")
        valid_statuses = {value for value, _ in PhysicalReturnReceipt.FollowUpStatus.choices}
        if follow_up_status not in valid_statuses:
            messages.error(request, "Status tindak lanjut tidak valid.")
        else:
            receipt.follow_up_status = follow_up_status
            receipt.save(update_fields=["follow_up_status"])
            messages.success(request, "Status tindak lanjut marketplace diperbarui.")
        return redirect(f"{request.get_full_path()}#marketplace-claims")

    if request.method == "POST":
        if not selected_order:
            form.add_error(None, "Pilih Source dan No. Pesanan terlebih dahulu.")
        entries = []
        for row in return_rows:
            field_name = f"quantity_{row['line'].id}"
            raw_quantity = request.POST.get(field_name, "").strip()
            submitted_quantities[str(row["line"].id)] = raw_quantity
            if not raw_quantity:
                continue
            try:
                quantity = int(raw_quantity)
            except ValueError:
                form.add_error(None, f"Qty Return {row['line'].sku.sku} harus berupa bilangan bulat.")
                continue
            if quantity <= 0 or Decimal(quantity) > row["remaining_qty"]:
                form.add_error(
                    None,
                    f"Qty Return {row['line'].sku.sku} harus 1–{row['remaining_qty']:.0f} pcs.",
                )
                continue
            entries.append((row["line"], quantity))
        if not entries:
            form.add_error(None, "Isi minimal satu Qty Return.")

        if form.is_valid() and selected_order and entries:
            restored_count = 0
            try:
                with transaction.atomic():
                    for sales_line, quantity in entries:
                        create_expected_return(sales_line)
                        _, movement = record_physical_return(
                            sales_line=sales_line,
                            quantity=quantity,
                            actor=request.user,
                            **form.cleaned_data,
                        )
                        restored_count += int(movement is not None)
            except ValidationError as exc:
                form.add_error(None, exc)
            else:
                messages.success(
                    request,
                    f"Sales Return diterima untuk {len(entries)} SKU; "
                    f"{restored_count} SKU memulihkan stock/FIFO.",
                )
                query = urlencode({"source": selected_source, "order_number": selected_order.order_number})
                return redirect(f"{reverse('inventory:return_log')}?{query}")

    receipt_rows = PhysicalReturnReceipt.objects.select_related(
        "sales_line__order",
        "sales_line__sku__product_variant__product",
        "warehouse",
        "recorded_by",
        "movement",
    ).order_by("-received_date", "-created_at")
    if request.method == "GET" and request.GET.get("export") == "xlsx":
        return _export_returns(
            return_rows,
            receipt_rows.filter(condition__in=claim_conditions),
            receipt_rows,
            selected_order,
        )
    receipts = receipt_rows[:300]
    claim_receipts = receipt_rows.filter(condition__in=claim_conditions)[:300]
    return render(
        request,
        "inventory/returns.html",
        {
            "form": form,
            "source_options": source_options,
            "selected_source": selected_source,
            "order_options": filtered_orders,
            "selected_order": selected_order,
            "return_rows": return_rows,
            "submitted_quantities": submitted_quantities,
            "receipts": receipts,
            "claim_receipts": claim_receipts,
            "follow_up_status_options": PhysicalReturnReceipt.FollowUpStatus.choices,
        },
    )


@login_required
def outbound(request):
    query = request.GET.get("q", "").strip()
    date_from = request.GET.get("date_from") or None
    date_to = request.GET.get("date_to") or None
    source = request.GET.get("source", "")
    rows = InventoryMovement.objects.filter(movement_type=InventoryMovement.MovementType.SALES_OUT).select_related(
        "sku", "sales_line__order", "posted_by"
    )
    if query:
        rows = rows.filter(Q(sku__sku__icontains=query) | Q(sales_line__order__order_number__icontains=query))
    if date_from:
        rows = rows.filter(movement_date__gte=date_from)
    if date_to:
        rows = rows.filter(movement_date__lte=date_to)
    if source:
        rows = rows.filter(sales_line__order__source=source)
    rows = rows.order_by("-movement_date", "-posted_at")
    if request.GET.get("export") == "xlsx":
        return _export_outbound(rows)
    page = Paginator(rows, 100).get_page(request.GET.get("page"))
    return render(request, "inventory/outbound.html", {"page": page, "query": query, "date_from": date_from, "date_to": date_to, "source": source})


@login_required
def opening_import_list(request):
    return render(request, "inventory/opening/list.html", {"batches": FIFOOpeningImportBatch.objects.select_related("raw_file", "approved_by")[:20]})


@login_required
def opening_import_upload(request):
    if request.method == "POST":
        form = FIFOOpeningImportUploadForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                batch = create_opening_import(form.cleaned_data["file"], request.user)
            except DuplicateRawFile as exc:
                existing = exc.raw_file.fifo_opening_batches.order_by("-created_at").first()
                form.add_error("file", f"File identik sudah pernah diunggah. Batch sebelumnya: {existing.id if existing else '-'}.")
            else:
                if batch.status == FIFOOpeningImportBatch.Status.READY:
                    messages.success(request, "FIFO Opening berhasil diparsing dan siap direview.")
                else:
                    messages.warning(request, "FIFO Opening diparsing tetapi memiliki blocking issue.")
                return redirect("inventory:opening_detail", batch_id=batch.id)
    elif request.method == "GET":
        form = FIFOOpeningImportUploadForm()
    else:
        return HttpResponseNotAllowed(["GET", "POST"])
    return render(request, "inventory/opening/upload.html", {"form": form})


@login_required
def opening_import_detail(request, batch_id):
    batch = get_object_or_404(FIFOOpeningImportBatch.objects.select_related("raw_file", "approved_by"), pk=batch_id)
    page = Paginator(batch.staged_rows.select_related("sku"), 50).get_page(request.GET.get("page"))
    issues = batch.issues.select_related("staged_row")
    severity = request.GET.get("severity", "")
    if severity in FIFOOpeningImportIssue.Severity.values:
        issues = issues.filter(severity=severity)
    return render(request, "inventory/opening/detail.html", {"batch": batch, "page": page, "issues": issues[:100]})


@login_required
def opening_import_approve(request, batch_id):
    if request.method != "POST":
        return HttpResponseNotAllowed(["POST"])
    batch = get_object_or_404(FIFOOpeningImportBatch, pk=batch_id)
    try:
        _, counts = approve_opening_import(batch.id, request.user)
    except ValidationError as exc:
        messages.error(request, " ".join(exc.messages))
    else:
        messages.success(request, f"FIFO Opening committed: {counts['snapshots']} snapshot, {counts['positive_layers']} layer positif, {counts['negative_exceptions']} exception negatif.")
    return redirect("inventory:opening_detail", batch_id=batch.id)
