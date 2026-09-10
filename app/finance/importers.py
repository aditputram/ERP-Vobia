import hashlib
import re
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import load_workbook

from audit.services import record_audit

from .models import Account, FINANCE_OPENING_DATE, JournalEntry, JournalLine


ZERO = Decimal("0")
TOLERANCE = Decimal("0.01")


@dataclass(frozen=True)
class SourceAccount:
    code: str
    name: str
    account_type: str
    parent_code: str
    currency: str


def _text(value):
    return str(value or "").strip()


def _amount(value):
    if value in (None, ""):
        return ZERO
    return Decimal(str(value))


def _normalized(value):
    return re.sub(r"\s+", " ", _text(value)).casefold()


def _sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_chart_of_accounts(path):
    sheet = load_workbook(path, data_only=True).active
    accounts = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        code = _text(row[2] if len(row) > 2 else "")
        if not code:
            continue
        accounts.append(
            SourceAccount(
                code=code,
                name=_text(row[3]),
                account_type=_text(row[1]),
                parent_code=_text(row[4]),
                currency=_text(row[5]) or "IDR",
            )
        )
    if not accounts:
        raise ValidationError("File Chart of Accounts tidak memiliki baris akun.")
    codes = [account.code for account in accounts]
    if len(codes) != len(set(codes)):
        raise ValidationError("File Chart of Accounts memiliki kode duplikat.")
    missing_parents = sorted(
        {account.parent_code for account in accounts if account.parent_code and account.parent_code not in set(codes)}
    )
    if missing_parents:
        raise ValidationError("Akun induk tidak ditemukan: " + ", ".join(missing_parents))
    return accounts


def parse_trial_balance(path, source_accounts):
    sheet = load_workbook(path, data_only=True).active
    by_code = {account.code: account for account in source_accounts}
    parent_codes = {account.parent_code for account in source_accounts if account.parent_code}
    report_rows = {}
    for row_number in range(6, sheet.max_row + 1):
        code = _text(sheet.cell(row_number, 2).value)
        if not code or not code.isdigit():
            continue
        name = _text(sheet.cell(row_number, 4).value)
        if code not in by_code:
            raise ValidationError(f"Kode Trial Balance {code} tidak ada di Chart of Accounts.")
        if _normalized(name) != _normalized(by_code[code].name):
            raise ValidationError(f"Nama akun {code} berbeda antara Trial Balance dan Chart of Accounts.")
        opening_debit = _amount(sheet.cell(row_number, 6).value)
        opening_credit = _amount(sheet.cell(row_number, 8).value)
        movement_debit = _amount(sheet.cell(row_number, 10).value)
        movement_credit = _amount(sheet.cell(row_number, 12).value)
        ending_debit = _amount(sheet.cell(row_number, 14).value)
        ending_credit = _amount(sheet.cell(row_number, 16).value)
        rollforward = opening_debit - opening_credit + movement_debit - movement_credit
        ending = ending_debit - ending_credit
        if abs(rollforward - ending) > TOLERANCE:
            raise ValidationError(f"Roll-forward Trial Balance tidak cocok pada akun {code}.")
        if opening_debit and opening_credit or ending_debit and ending_credit:
            raise ValidationError(f"Akun {code} terisi pada sisi Debit dan Kredit sekaligus.")
        report_rows[code] = {"debit": ending_debit, "credit": ending_credit}

    leaf_rows = {
        code: amount
        for code, amount in report_rows.items()
        if code not in parent_codes and (amount["debit"] or amount["credit"])
    }
    debit_total = sum((amount["debit"] for amount in leaf_rows.values()), ZERO)
    credit_total = sum((amount["credit"] for amount in leaf_rows.values()), ZERO)
    if debit_total <= 0 or abs(debit_total - credit_total) > TOLERANCE:
        raise ValidationError(
            f"Trial Balance tidak balance. Debit {debit_total} dan Kredit {credit_total}."
        )
    return leaf_rows, debit_total, credit_total


@transaction.atomic
def stage_finance_cutover(*, coa_path, trial_balance_path, cutoff_date, actor=None, replace_staged=False):
    coa_path = Path(coa_path)
    trial_balance_path = Path(trial_balance_path)
    accounts = parse_chart_of_accounts(coa_path)
    opening_rows, debit_total, credit_total = parse_trial_balance(trial_balance_path, accounts)

    parent_codes = {account.parent_code for account in accounts if account.parent_code}
    for source in accounts:
        Account.objects.update_or_create(
            code=source.code,
            defaults={
                "name": source.name,
                "account_type": source.account_type,
                "currency": source.currency,
                "is_postable": source.code not in parent_codes,
                "is_active": True,
            },
        )
    account_map = {account.code: account for account in Account.objects.all()}
    for source in accounts:
        parent = account_map.get(source.parent_code) if source.parent_code else None
        account = account_map[source.code]
        if account.parent_id != getattr(parent, "id", None):
            account.parent = parent
            account.save(update_fields=("parent", "updated_at"))

    number = f"OPENING-{cutoff_date:%Y%m%d}"
    existing = JournalEntry.objects.filter(number=number).first()
    if existing:
        if existing.status == JournalEntry.Status.POSTED:
            raise ValidationError("Opening journal sudah Posted dan tidak boleh diganti.")
        if not replace_staged:
            raise ValidationError("Opening journal Draft sudah ada. Gunakan --replace-staged untuk membangun ulang.")
        existing.lines.all().delete()
        entry = existing
        entry.entry_date = FINANCE_OPENING_DATE
        entry.description = f"Saldo awal berdasarkan Trial Balance per {cutoff_date:%d %B %Y}"
        entry.reference = f"TB-{cutoff_date:%Y-%m-%d}"
        entry.created_by = actor
    else:
        entry = JournalEntry(
            number=number,
            entry_date=FINANCE_OPENING_DATE,
            description=f"Saldo awal berdasarkan Trial Balance per {cutoff_date:%d %B %Y}",
            reference=f"TB-{cutoff_date:%Y-%m-%d}",
            source=JournalEntry.Source.OPENING,
            created_by=actor,
        )
    entry.source_metadata = {
        "cutoff_date": cutoff_date.isoformat(),
        "coa_filename": coa_path.name,
        "coa_sha256": _sha256(coa_path),
        "trial_balance_filename": trial_balance_path.name,
        "trial_balance_sha256": _sha256(trial_balance_path),
        "reconciliation_status": "PENDING_CONTROL_ACCOUNT_RECONCILIATION",
    }
    entry.full_clean()
    entry.save()
    JournalLine.objects.bulk_create(
        [
            JournalLine(
                entry=entry,
                line_number=index,
                account=account_map[code],
                description="Opening 1 September 2026",
                debit=amount["debit"],
                credit=amount["credit"],
            )
            for index, (code, amount) in enumerate(sorted(opening_rows.items()), start=1)
        ]
    )
    record_audit(
        actor=actor,
        action="finance_cutover_staged",
        entity_type="finance.journal_entry",
        entity_id=entry.id,
        after_values={
            "cutoff_date": cutoff_date.isoformat(),
            "accounts": len(accounts),
            "opening_lines": len(opening_rows),
            "debit": str(debit_total),
            "credit": str(credit_total),
        },
        metadata=entry.source_metadata,
    )
    return entry
