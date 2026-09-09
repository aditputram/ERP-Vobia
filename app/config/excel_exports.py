from io import BytesIO

from django.http import HttpResponse
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def safe_cell(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def append_table(sheet, headers, rows, *, number_formats=None):
    sheet.sheet_view.showGridLines = False
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="151713")
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"

    for values in rows:
        sheet.append(tuple(safe_cell(value) for value in values))

    if sheet.max_row:
        sheet.auto_filter.ref = sheet.dimensions
    for column, number_format in (number_formats or {}).items():
        for row in range(2, sheet.max_row + 1):
            sheet.cell(row=row, column=column).number_format = number_format
    for column in range(1, sheet.max_column + 1):
        width = max(
            len(str(sheet.cell(row=row, column=column).value or ""))
            for row in range(1, min(sheet.max_row, 250) + 1)
        )
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 12), 36)


def workbook_response(workbook, filename):
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    response = HttpResponse(output.getvalue(), content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
