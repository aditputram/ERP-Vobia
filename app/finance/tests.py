import uuid
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook

from audit.models import AuditEvent
from master_data.models import Category, Product, ProductStatus, ProductVariant, SKU, Subcategory

from .importers import stage_finance_cutover
from .catalog import FEATURES
from .models import (
    Account,
    JournalEntry,
    JournalLine,
    ProductSalesAccount,
    SalesJournalAllocation,
    SalesReturnJournalAllocation,
)
from .services import (
    account_balances,
    create_sales_journal_draft,
    create_sales_return_journal_draft,
    finance_sales_lines,
    post_journal,
    sales_return_journal_preview,
)


class FinanceJournalTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_superuser(username="finance-admin", password="test")
        self.cash, _ = Account.objects.update_or_create(
            code="110101", defaults={"name": "Cash", "account_type": "BANK", "is_postable": True}
        )
        self.capital, _ = Account.objects.update_or_create(
            code="300001", defaults={"name": "Capital", "account_type": "EQTY", "is_postable": True}
        )
        self.entry = JournalEntry.objects.create(
            number="JV-202609-0001",
            entry_date=date(2026, 9, 1),
            description="Setoran modal",
            created_by=self.user,
        )
        JournalLine.objects.create(entry=self.entry, line_number=1, account=self.cash, debit=100, credit=0)
        JournalLine.objects.create(entry=self.entry, line_number=2, account=self.capital, debit=0, credit=100)

    def _sales_sku(self, code, sales_account_code="410002"):
        status, _ = ProductStatus.objects.get_or_create(code="FIN-JOURNAL", defaults={"name": "Finance Journal"})
        category, _ = Category.objects.get_or_create(code="FIN-JOURNAL", defaults={"name": "Finance Journal"})
        product = Product.objects.create(code=code, name=code, status=status, category=category)
        variant = ProductVariant.objects.create(product=product, name=code)
        sku = SKU.objects.create(sku=code, product_variant=variant)
        ProductSalesAccount.objects.create(
            product=product,
            sales_account=Account.objects.get(code=sales_account_code),
        )
        return sku

    def _allocate_sales_line(self, sales_line, receipt_account_code="110301"):
        amount = sales_line.total_net_sales or (sales_line.quantity * sales_line.net_unit_price)
        entry = JournalEntry.objects.create(
            number=f"JV-SALES-{uuid.uuid4().hex[:12].upper()}",
            entry_date=sales_line.order.order_date,
            description="Original Sales journal",
            source=JournalEntry.Source.SYSTEM,
            source_metadata={"workflow": "SALES_JOURNAL_BATCH"},
            created_by=self.user,
        )
        JournalLine.objects.create(
            entry=entry,
            line_number=1,
            account=Account.objects.get(code=receipt_account_code),
            debit=amount,
            credit=0,
            description="Sales Receivable / Receipt",
        )
        JournalLine.objects.create(
            entry=entry,
            line_number=2,
            account=Account.objects.get(code="410002"),
            debit=0,
            credit=amount,
            description="Gross Sales",
        )
        SalesJournalAllocation.objects.create(
            entry=entry,
            sales_line=sales_line,
            gross_sales=amount,
            net_sales=amount,
            cogs=0,
        )
        return entry

    def test_balanced_journal_posts_and_updates_trial_balance(self):
        post_journal(self.entry.id, self.user)
        self.entry.refresh_from_db()
        self.assertEqual(self.entry.status, JournalEntry.Status.POSTED)
        balances = {row["account"].code: row for row in account_balances(end_date=date(2026, 9, 1))}
        self.assertEqual(balances["110101"]["debit_balance"], Decimal("100"))
        self.assertEqual(balances["300001"]["credit_balance"], Decimal("100"))
        self.assertTrue(AuditEvent.objects.filter(action="finance_journal_posted").exists())

    def test_unbalanced_journal_cannot_be_posted(self):
        self.entry.lines.filter(line_number=2).update(credit=90)
        with self.assertRaises(ValidationError):
            post_journal(self.entry.id, self.user)

    def test_opening_journal_cannot_post_before_control_reconciliation(self):
        self.entry.source = JournalEntry.Source.OPENING
        self.entry.source_metadata = {"reconciliation_status": "PENDING_CONTROL_ACCOUNT_RECONCILIATION"}
        self.entry.save(update_fields=("source", "source_metadata"))

        with self.assertRaisesMessage(ValidationError, "akun kontrol"):
            post_journal(self.entry.id, self.user)

    def test_period_balance_can_exclude_opening_journal(self):
        self.entry.status = JournalEntry.Status.POSTED
        self.entry.save(update_fields=("status",))
        opening = JournalEntry.objects.create(
            number="OPENING-TEST",
            entry_date=date(2026, 9, 1),
            description="Opening",
            source=JournalEntry.Source.OPENING,
            status=JournalEntry.Status.POSTED,
        )
        JournalLine.objects.create(entry=opening, line_number=1, account=self.cash, debit=50, credit=0)
        JournalLine.objects.create(entry=opening, line_number=2, account=self.capital, debit=0, credit=50)

        balances = {
            row["account"].code: row
            for row in account_balances(
                start_date=date(2026, 9, 1),
                end_date=date(2026, 9, 30),
                exclude_opening=True,
            )
        }
        self.assertEqual(balances["110101"]["debit_balance"], Decimal("100"))

    def test_finance_pages_render(self):
        self.client.force_login(self.user)
        for name in ("dashboard", "accounts", "journals", "trial_balance", "balance_sheet", "profit_loss"):
            with self.subTest(name=name):
                self.assertEqual(self.client.get(reverse(f"finance:{name}")).status_code, 200)

        self.assertContains(self.client.get(reverse("finance:dashboard")), "Finance UAT")

    def test_general_ledger_can_show_all_accounts_and_selected_account_detail(self):
        self.client.force_login(self.user)
        url = reverse("finance:feature", args=("general-ledger-summary",))

        summary = self.client.get(url, {"start": "2026-09-01", "end": "2026-09-30"})
        self.assertContains(summary, "Semua Akun")
        self.assertContains(summary, "finance-ledger-filter")
        self.assertContains(summary, 'type="search"')
        self.assertContains(summary, 'list="ledger-account-options"')
        self.assertContains(summary, "data-ledger-account-value")
        self.assertContains(summary, "110101")

        detail = self.client.get(
            url,
            {"start": "2026-09-01", "end": "2026-09-30", "account": self.cash.id},
        )
        self.assertContains(detail, self.entry.number)
        self.assertContains(detail, "Draft")
        self.assertContains(detail, "Saldo awal")
        self.assertContains(detail, "Rp 100,00 D")
        self.assertContains(detail, f"{self.cash.code} · {self.cash.name}")

    def test_superadmin_can_add_coa_from_account_page(self):
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("finance:accounts"),
            {
                "code": "990001",
                "name": "UAT Finance Account",
                "account_type": "OEXP",
                "currency": "IDR",
                "is_postable": "on",
                "is_active": "on",
            },
        )

        self.assertRedirects(response, reverse("finance:accounts"))
        self.assertTrue(Account.objects.filter(code="990001", name="UAT Finance Account").exists())
        self.assertTrue(AuditEvent.objects.filter(action="finance_account_created").exists())

    def test_view_only_user_cannot_add_coa(self):
        user = get_user_model().objects.create_user(
            username="finance-viewer",
            password="test",
            module_access={"finance": "view"},
            tab_access={"finance": ["accounts"]},
        )
        self.client.force_login(user)

        page = self.client.get(reverse("finance:accounts"))

        self.assertEqual(page.status_code, 200)
        self.assertNotContains(page, "Tambah COA")
        self.assertEqual(self.client.post(reverse("finance:accounts"), {}).status_code, 403)

    def test_inventory_account_is_included_as_balance_sheet_asset(self):
        inventory, _ = Account.objects.update_or_create(
            code="110401", defaults={"name": "Inventory", "account_type": "INTR", "is_postable": True}
        )
        entry = JournalEntry.objects.create(
            number="JV-202609-0002",
            entry_date=date(2026, 9, 1),
            description="Inventory opening",
            status=JournalEntry.Status.POSTED,
        )
        JournalLine.objects.create(entry=entry, line_number=1, account=inventory, debit=50, credit=0)
        JournalLine.objects.create(entry=entry, line_number=2, account=self.capital, debit=0, credit=50)
        self.client.force_login(self.user)

        response = self.client.get(reverse("finance:balance_sheet"), {"as_of": "2026-09-01"})

        self.assertContains(response, "Inventory")
        self.assertEqual(response.context["difference"], Decimal("0"))

    def test_balance_sheet_uses_draft_opening_but_excludes_other_drafts(self):
        inventory, _ = Account.objects.update_or_create(
            code="110401",
            defaults={"name": "Opening Inventory", "account_type": "INTR", "is_postable": True},
        )
        opening = JournalEntry.objects.get(number="OPENING-20260831")
        opening.lines.all().delete()
        opening.status = JournalEntry.Status.DRAFT
        opening.save(update_fields=("status",))
        JournalLine.objects.create(entry=opening, line_number=1, account=inventory, debit=50, credit=0)
        JournalLine.objects.create(entry=opening, line_number=2, account=self.capital, debit=0, credit=50)
        self.client.force_login(self.user)

        response = self.client.get(
            reverse("finance:balance_sheet"),
            {"as_of": "2026-09-22", "mode": "posted"},
        )

        self.assertContains(response, "Opening Inventory")
        self.assertEqual(response.context["totals"]["assets"], Decimal("50"))
        self.assertEqual(response.context["totals"]["equity"], Decimal("50"))
        self.assertEqual(response.context["difference"], Decimal("0"))

    def test_regular_user_does_not_gain_finance_access_by_default(self):
        user = get_user_model().objects.create_user(username="legacy", password="test")
        self.client.force_login(user)
        self.assertEqual(self.client.get(reverse("finance:dashboard")).status_code, 403)

    def test_every_requested_finance_workspace_renders(self):
        self.client.force_login(self.user)
        for slug in FEATURES:
            with self.subTest(slug=slug):
                self.assertEqual(self.client.get(reverse("finance:feature", args=[slug])).status_code, 200)

    def test_sales_invoice_uses_canonical_sales_transactions(self):
        from sales.models import SalesOrder, SalesOrderLine

        order = SalesOrder.objects.create(
            source=SalesOrder.Source.SHOPEE,
            source_label="Shopee",
            order_number="FINANCE-SALES-001",
            order_datetime=timezone.make_aware(datetime(2026, 9, 10, 10, 0)),
            order_date=date(2026, 9, 10),
            current_status="Selesai",
            source_status="Selesai",
            is_final=True,
            first_seen_batch_id=uuid.uuid4(),
            latest_batch_id=uuid.uuid4(),
        )
        SalesOrderLine.objects.create(
            order=order,
            sku_code_snapshot="FINANCE-SKU-001",
            product_name_snapshot="Finance Sales Product",
            quantity=2,
            net_unit_price=Decimal("90000"),
            retail_price_snapshot=Decimal("100000"),
            total_gross_sales=Decimal("200000"),
            total_net_sales=Decimal("180000"),
            total_cogs=Decimal("120000"),
        )
        older_order = SalesOrder.objects.create(
            source=SalesOrder.Source.TIKTOK,
            source_label="TikTok",
            order_number="FINANCE-SALES-OLDER",
            order_datetime=timezone.make_aware(datetime(2026, 8, 10, 10, 0)),
            order_date=date(2026, 8, 10),
            current_status="Selesai",
            source_status="Selesai",
            is_final=True,
            first_seen_batch_id=uuid.uuid4(),
            latest_batch_id=uuid.uuid4(),
        )
        SalesOrderLine.objects.create(
            order=older_order,
            sku_code_snapshot="FINANCE-SKU-OLDER",
            product_name_snapshot="Older Finance Product",
            quantity=1,
            net_unit_price=Decimal("75000"),
            retail_price_snapshot=Decimal("100000"),
            total_gross_sales=Decimal("100000"),
            total_net_sales=Decimal("75000"),
            total_cogs=Decimal("50000"),
        )
        self.client.force_login(self.user)

        response = self.client.get(
            reverse("finance:feature", args=["sales-invoice"]),
            {"period_type": "month", "period": "2026-09", "source": "Shopee"},
        )

        self.assertEqual(response.context["sales_totals"]["orders"], 1)
        self.assertEqual(response.context["sales_totals"]["cogs"], Decimal("120000"))
        self.assertEqual(response.context["selected_sources"], ["Shopee"])
        self.assertContains(response, "FINANCE-SALES-001")
        self.assertNotContains(response, "FINANCE-SALES-OLDER")
        self.assertContains(response, "Shopee")
        self.assertContains(response, "200.000")
        self.assertContains(response, "180.000")
        self.assertContains(response, "120.000")
        self.assertContains(response, "COGS")
        self.assertNotContains(response, "Buat Sales Invoice")

    def test_sales_invoice_waits_until_order_is_shipped(self):
        from sales.models import SalesOrder, SalesOrderLine

        sku = self._sales_sku("FINANCE-PENDING-SKU")
        order = SalesOrder.objects.create(
            source=SalesOrder.Source.SHOPEE,
            source_label="Shopee",
            order_number="FINANCE-PENDING-001",
            order_datetime=timezone.make_aware(datetime(2026, 9, 10, 10, 0)),
            order_date=date(2026, 9, 10),
            current_status="Belum Dibayar",
            source_status="Belum Dibayar",
            first_seen_batch_id=uuid.uuid4(),
            latest_batch_id=uuid.uuid4(),
        )
        line = SalesOrderLine.objects.create(
            order=order,
            sku=sku,
            current_status="Belum Dibayar",
            quantity=1,
            net_unit_price=Decimal("90000"),
            retail_price_snapshot=Decimal("100000"),
            total_gross_sales=Decimal("100000"),
            total_net_sales=Decimal("90000"),
            total_cogs=Decimal("60000"),
        )
        self.client.force_login(self.user)

        pending_page = self.client.get(
            reverse("finance:feature", args=["sales-invoice"]),
            {"period_type": "month", "period": "2026-09"},
        )
        self.assertNotContains(pending_page, order.order_number)
        self.assertFalse(finance_sales_lines().filter(pk=line.pk).exists())

        line.current_status = "Perlu Dikirim"
        line.save(update_fields=("current_status",))
        order.shipped_datetime = timezone.make_aware(datetime(2026, 9, 10, 12, 0))
        order.save(update_fields=("shipped_datetime",))
        self.assertFalse(finance_sales_lines().filter(pk=line.pk).exists())
        with self.assertRaisesMessage(ValidationError, "belum dijurnal"):
            create_sales_journal_draft(
                start_date=date(2026, 9, 1),
                end_date=date(2026, 9, 30),
                receipt_account_id=Account.objects.get(code="110301").id,
                actor=self.user,
            )

        line.current_status = "Telah Dikirim"
        line.save(update_fields=("current_status",))
        order.current_status = "Telah Dikirim"
        order.source_status = "Telah Dikirim"
        order.save(update_fields=("current_status", "source_status"))

        shipped_page = self.client.get(
            reverse("finance:feature", args=["sales-invoice"]),
            {"period_type": "month", "period": "2026-09"},
        )
        self.assertContains(shipped_page, order.order_number)
        self.assertTrue(finance_sales_lines().filter(pk=line.pk).exists())

    def test_sales_invoice_filters_transaction_status(self):
        from sales.models import SalesOrder, SalesOrderLine

        for index, status in enumerate(("Dikirim", "Selesai"), 1):
            order = SalesOrder.objects.create(
                source=SalesOrder.Source.SHOPEE,
                source_label="Shopee",
                order_number=f"FINANCE-STATUS-{index}",
                order_datetime=timezone.make_aware(datetime(2026, 9, 10, 10, index)),
                order_date=date(2026, 9, 10),
                current_status=status,
                source_status=status,
                is_final=True,
                first_seen_batch_id=uuid.uuid4(),
                latest_batch_id=uuid.uuid4(),
            )
            SalesOrderLine.objects.create(
                order=order,
                sku_code_snapshot=f"FINANCE-STATUS-SKU-{index}",
                product_name_snapshot="Finance Status Product",
                quantity=1,
                net_unit_price=Decimal("90000"),
                retail_price_snapshot=Decimal("100000"),
                total_gross_sales=Decimal("100000"),
                total_net_sales=Decimal("90000"),
                total_cogs=Decimal("60000"),
            )
        self.client.force_login(self.user)

        response = self.client.get(
            reverse("finance:feature", args=["sales-invoice"]),
            {"period_type": "month", "period": "2026-09", "status": "Dikirim"},
        )

        self.assertEqual(response.context["selected_status"], "Dikirim")
        self.assertEqual(response.context["status_options"], ["Dikirim", "Selesai"])
        self.assertEqual(response.context["sales_totals"]["orders"], 1)
        self.assertContains(response, "FINANCE-STATUS-1")
        self.assertNotContains(response, "FINANCE-STATUS-2")

    def test_sales_journal_is_balanced_draft_and_does_not_duplicate_sales_lines(self):
        from sales.models import SalesOrder, SalesOrderLine

        sku = self._sales_sku("FIN-JOURNAL-SKU")
        order = SalesOrder.objects.create(
            source=SalesOrder.Source.SHOPEE,
            source_label="Shopee",
            order_number="FINANCE-JOURNAL-001",
            order_datetime=timezone.make_aware(datetime(2026, 9, 12, 10, 0)),
            order_date=date(2026, 9, 12),
            current_status="Selesai",
            source_status="Selesai",
            is_final=True,
            first_seen_batch_id=uuid.uuid4(),
            latest_batch_id=uuid.uuid4(),
        )
        sales_line = SalesOrderLine.objects.create(
            order=order,
            sku=sku,
            sku_code_snapshot="FIN-JOURNAL-SKU",
            category_snapshot="Finance T-Shirt",
            product_name_snapshot="Finance Journal Product",
            quantity=2,
            net_unit_price=Decimal("90000"),
            retail_price_snapshot=Decimal("100000"),
            total_gross_sales=Decimal("200000"),
            total_net_sales=Decimal("180000"),
            total_cogs=Decimal("120000"),
        )
        params = {
            "start_date": date(2026, 9, 1),
            "end_date": date(2026, 9, 30),
            "receipt_account_id": Account.objects.get(code="110301").id,
            "actor": self.user,
        }

        entry = create_sales_journal_draft(**params)

        self.assertEqual(entry.status, JournalEntry.Status.DRAFT)
        self.assertEqual(entry.debit_total, Decimal("320000"))
        self.assertEqual(entry.credit_total, Decimal("320000"))
        self.assertTrue(SalesJournalAllocation.objects.filter(sales_line=sales_line, entry=entry).exists())
        with self.assertRaisesMessage(ValidationError, "belum dijurnal"):
            create_sales_journal_draft(**params)

        self.client.force_login(self.user)
        page = self.client.get(reverse("finance:feature", args=["sales-invoice"]))
        self.assertContains(page, "Create Jurnal Entry Sales")
        self.assertContains(page, "1 baris sudah dialokasikan")
        self.assertContains(page, 'name="receipt_account"')
        self.assertNotContains(page, 'name="discount_account"')
        self.assertNotContains(page, 'name="inventory_account"')
        self.assertNotContains(page, 'name="sales_mode"')
        self.assertNotContains(page, 'name="cogs_mode"')
        dialog = page.content.decode().split('data-sales-journal-dialog', 1)[1]
        self.assertIn('name="source_group"', dialog)
        self.assertIn('name="source"', dialog)
        self.assertIn('name="category"', dialog)

        late_order = SalesOrder.objects.create(
            source=SalesOrder.Source.SHOPEE,
            source_label="Shopee",
            order_number="FINANCE-JOURNAL-LATE-001",
            order_datetime=timezone.make_aware(datetime(2026, 9, 12, 15, 0)),
            order_date=date(2026, 9, 12),
            current_status="Selesai",
            source_status="Selesai",
            is_final=True,
            first_seen_batch_id=uuid.uuid4(),
            latest_batch_id=uuid.uuid4(),
        )
        late_sales_line = SalesOrderLine.objects.create(
            order=late_order,
            sku=sku,
            sku_code_snapshot="FIN-JOURNAL-SKU",
            category_snapshot="Finance T-Shirt",
            product_name_snapshot="Finance Journal Product",
            quantity=1,
            net_unit_price=Decimal("90000"),
            retail_price_snapshot=Decimal("100000"),
            total_gross_sales=Decimal("100000"),
            total_net_sales=Decimal("90000"),
            total_cogs=Decimal("60000"),
        )

        late_entry = create_sales_journal_draft(**params)

        self.assertTrue(
            SalesJournalAllocation.objects.filter(
                sales_line=late_sales_line,
                entry=late_entry,
            ).exists()
        )
        self.assertEqual(
            SalesJournalAllocation.objects.filter(sales_line=sales_line).count(),
            1,
        )

        pants_sku = self._sales_sku("FIN-JOURNAL-PANTS", "410006")
        october_order = SalesOrder.objects.create(
            source=SalesOrder.Source.SHOPEE,
            source_label="Shopee",
            order_number="FINANCE-JOURNAL-002",
            order_datetime=timezone.make_aware(datetime(2026, 10, 1, 10, 0)),
            order_date=date(2026, 10, 1),
            current_status="Selesai",
            source_status="Selesai",
            is_final=True,
            first_seen_batch_id=uuid.uuid4(),
            latest_batch_id=uuid.uuid4(),
        )
        october_shirt = SalesOrderLine.objects.create(
            order=october_order,
            sku=sku,
            category_snapshot="Finance T-Shirt",
            quantity=1,
            net_unit_price=Decimal("90000"),
            retail_price_snapshot=Decimal("100000"),
            total_gross_sales=Decimal("100000"),
            total_net_sales=Decimal("90000"),
            total_cogs=Decimal("60000"),
        )
        october_pants_order = SalesOrder.objects.create(
            source=SalesOrder.Source.TIKTOK,
            source_label="TikTok",
            order_number="FINANCE-JOURNAL-003",
            order_datetime=timezone.make_aware(datetime(2026, 10, 1, 11, 0)),
            order_date=date(2026, 10, 1),
            current_status="Selesai",
            source_status="Selesai",
            is_final=True,
            first_seen_batch_id=uuid.uuid4(),
            latest_batch_id=uuid.uuid4(),
        )
        october_pants = SalesOrderLine.objects.create(
            order=october_pants_order,
            sku=pants_sku,
            category_snapshot="Finance Pants",
            quantity=1,
            net_unit_price=Decimal("180000"),
            retail_price_snapshot=Decimal("200000"),
            total_gross_sales=Decimal("200000"),
            total_net_sales=Decimal("180000"),
            total_cogs=Decimal("120000"),
        )
        response = self.client.post(
            reverse("finance:feature", args=["sales-invoice"]),
            {
                "action": "create_sales_journal",
                "journal_start": "2026-10-01",
                "journal_end": "2026-10-01",
                "receipt_account": Account.objects.get(code="110301").id,
                "source_group": ["Marketplace"],
                "source": ["Shopee"],
                "category": ["Finance T-Shirt"],
            },
        )
        self.assertEqual(response.status_code, 302)
        filtered_entry = JournalEntry.objects.get(reference="Sales 2026-10-01/2026-10-01")
        self.assertTrue(SalesJournalAllocation.objects.filter(entry=filtered_entry, sales_line=october_shirt).exists())
        self.assertFalse(SalesJournalAllocation.objects.filter(sales_line=october_pants).exists())
        self.assertEqual(filtered_entry.source_metadata["sources"], ["Shopee"])
        self.assertEqual(filtered_entry.source_metadata["categories"], ["Finance T-Shirt"])

    def test_sales_return_only_shows_returns_received_by_warehouse(self):
        from inventory.models import PhysicalReturnReceipt
        from master_data.models import Warehouse
        from sales.models import SalesOrder, SalesOrderLine

        order = SalesOrder.objects.create(
            source=SalesOrder.Source.SHOPEE,
            source_label="Shopee",
            order_number="FINANCE-RETURN-001",
            order_datetime=timezone.make_aware(datetime(2026, 8, 12, 10, 0)),
            order_date=date(2026, 8, 12),
            current_status="Retur",
            source_status="Retur",
            is_final=True,
            first_seen_batch_id=uuid.uuid4(),
            latest_batch_id=uuid.uuid4(),
        )
        received_line = SalesOrderLine.objects.create(
            order=order,
            sku_code_snapshot="FINANCE-RETURN-RECEIVED",
            product_name_snapshot="Received Return Product",
            current_status="Retur",
            quantity=2,
            net_unit_price=Decimal("90000"),
            total_net_sales=Decimal("180000"),
        )
        damaged_line = SalesOrderLine.objects.create(
            order=order,
            sku_code_snapshot="FINANCE-RETURN-DAMAGED",
            product_name_snapshot="Damaged Return Product",
            current_status="Retur",
            quantity=1,
            net_unit_price=Decimal("80000"),
            total_net_sales=Decimal("80000"),
        )
        SalesOrderLine.objects.create(
            order=order,
            sku_code_snapshot="FINANCE-RETURN-PENDING",
            product_name_snapshot="Pending Return Product",
            current_status="Retur",
            quantity=1,
            net_unit_price=Decimal("70000"),
            total_net_sales=Decimal("70000"),
        )
        warehouse = Warehouse.objects.create(code="FIN-RET-WH", name="Finance Return Warehouse")
        PhysicalReturnReceipt.objects.create(
            sales_line=received_line,
            received_date=date(2026, 9, 15),
            quantity=2,
            warehouse=warehouse,
            condition=PhysicalReturnReceipt.Condition.SELLABLE,
            recorded_by=self.user,
        )
        damaged_receipt = PhysicalReturnReceipt.objects.create(
            sales_line=damaged_line,
            received_date=date(2026, 9, 16),
            quantity=1,
            warehouse=warehouse,
            condition=PhysicalReturnReceipt.Condition.DAMAGED,
            recorded_by=self.user,
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse("finance:feature", args=["sales-return"]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "FINANCE-RETURN-RECEIVED")
        self.assertContains(response, "Received Return Product")
        self.assertContains(response, "Sellable")
        self.assertContains(response, "Damaged Return Product")
        self.assertContains(response, "Damaged")
        self.assertContains(response, "Finance Return Warehouse")
        self.assertNotContains(response, "FINANCE-RETURN-PENDING")
        self.assertNotContains(response, "Buat Sales Return")
        self.assertContains(response, "Create Jurnal Entry Sales Return")
        return_dialog = response.content.decode().split("data-return-journal-form", 1)[1]
        self.assertNotIn('name="receipt_account"', return_dialog)
        self.assertIn("PREVIEW JURNAL", return_dialog)
        self.assertEqual(response.context["rows"][1][5], 2)
        self.assertNotContains(response, "2,0000")
        self.assertEqual(
            response.context["metrics"],
            (("Return received", 2), ("Qty received", Decimal("3")), ("Belum dijurnal", 2)),
        )
        self.assertEqual(
            [value for value, _label in response.context["journal_receipt_status_options"]],
            [PhysicalReturnReceipt.Condition.SELLABLE, PhysicalReturnReceipt.Condition.DAMAGED],
        )
        self.assertEqual(
            {item["condition"] for item in response.context["journal_condition_filter"]["events"]},
            {PhysicalReturnReceipt.Condition.SELLABLE, PhysicalReturnReceipt.Condition.DAMAGED},
        )

        preview_response = self.client.get(
            reverse("finance:feature", args=["sales-return"]),
            {
                "journal_preview": "1",
                "journal_start": "2026-09-16",
                "journal_end": "2026-09-16",
                "journal_condition": PhysicalReturnReceipt.Condition.DAMAGED,
            },
        )
        self.assertEqual(preview_response.status_code, 200)
        self.assertEqual(preview_response.json()["receipt_count"], 1)
        self.assertEqual(preview_response.json()["receipt_quantity"], 1)
        self.assertEqual(preview_response.json()["transaction_count"], 1)
        self.assertEqual(preview_response.json()["lines"][1]["account_code"], "110301")

        narrow_response = self.client.post(
            reverse("finance:feature", args=["sales-return"]),
            {
                "action": "create_sales_return_journal",
                "journal_start": "2026-09-15",
                "journal_end": "2026-09-15",
            },
        )
        self.assertEqual(narrow_response.status_code, 200)
        self.assertEqual(
            [value for value, _label in narrow_response.context["journal_receipt_status_options"]],
            [PhysicalReturnReceipt.Condition.SELLABLE],
        )
        self.assertTrue(narrow_response.context["open_sales_return_journal_modal"])

        sellable_response = self.client.get(
            reverse("finance:feature", args=["sales-return"]),
            {"receipt_status": PhysicalReturnReceipt.Condition.SELLABLE},
        )
        self.assertContains(sellable_response, "Received Return Product")
        self.assertNotContains(sellable_response, "Damaged Return Product")
        self.assertEqual(
            sellable_response.context["metrics"],
            (("Return received", 1), ("Qty received", Decimal("2")), ("Belum dijurnal", 1)),
        )

        search_response = self.client.get(
            reverse("finance:feature", args=["sales-return"]),
            {"q": "FINANCE-RETURN-DAMAGED"},
        )
        self.assertContains(search_response, "Damaged Return Product")
        self.assertNotContains(search_response, "Received Return Product")
        self.assertEqual(search_response.context["query"], "FINANCE-RETURN-DAMAGED")

        month_response = self.client.get(
            reverse("finance:feature", args=["sales-return"]),
            {"period_type": "month", "period": "2026-09"},
        )
        self.assertEqual(month_response.context["period_type"], "month")
        self.assertEqual(month_response.context["period_value"], "2026-09")
        self.assertContains(month_response, "Received Return Product")
        self.assertContains(month_response, "Damaged Return Product")

        custom_response = self.client.get(
            reverse("finance:feature", args=["sales-return"]),
            {
                "period_type": "custom",
                "date_from": "2026-09-16",
                "date_to": "2026-09-16",
            },
        )
        self.assertNotContains(custom_response, "Received Return Product")
        self.assertContains(custom_response, "Damaged Return Product")
        self.assertEqual(
            custom_response.context["metrics"],
            (("Return received", 1), ("Qty received", Decimal("1")), ("Belum dijurnal", 1)),
        )

        create_response = self.client.post(
            reverse("finance:feature", args=["sales-return"]),
            {
                "action": "create_sales_return_journal",
                "journal_start": "2026-09-01",
                "journal_end": "2026-09-30",
                "journal_condition": PhysicalReturnReceipt.Condition.DAMAGED,
            },
        )
        self.assertEqual(create_response.status_code, 302)
        self.assertTrue(
            SalesReturnJournalAllocation.objects.filter(return_receipt=damaged_receipt).exists()
        )

    def test_sales_return_journal_reverses_revenue_and_sellable_cogs_without_duplicates(self):
        from inventory.models import InventoryMovement, PhysicalReturnReceipt
        from master_data.models import Warehouse
        from sales.models import SalesOrder, SalesOrderLine

        sellable_sku = self._sales_sku("FIN-RETURN-JOURNAL-SELLABLE")
        damaged_sku = self._sales_sku("FIN-RETURN-JOURNAL-DAMAGED")
        order = SalesOrder.objects.create(
            source=SalesOrder.Source.SHOPEE,
            source_label="Shopee",
            order_number="FINANCE-RETURN-JOURNAL-001",
            order_datetime=timezone.make_aware(datetime(2026, 9, 12, 10, 0)),
            order_date=date(2026, 9, 12),
            current_status="Retur",
            source_status="Retur",
            is_final=True,
            first_seen_batch_id=uuid.uuid4(),
            latest_batch_id=uuid.uuid4(),
        )
        sellable_line = SalesOrderLine.objects.create(
            order=order,
            sku=sellable_sku,
            product_name_snapshot="Sellable Return",
            current_status="Retur",
            quantity=2,
            net_unit_price=Decimal("90000"),
            total_net_sales=Decimal("180000"),
        )
        damaged_line = SalesOrderLine.objects.create(
            order=order,
            sku=damaged_sku,
            product_name_snapshot="Damaged Return",
            current_status="Retur",
            quantity=1,
            net_unit_price=Decimal("80000"),
            total_net_sales=Decimal("80000"),
        )
        warehouse = Warehouse.objects.create(code="FIN-RET-JOURNAL-WH", name="Return Journal Warehouse")
        sellable_receipt = PhysicalReturnReceipt.objects.create(
            sales_line=sellable_line,
            received_date=date(2026, 9, 20),
            quantity=2,
            warehouse=warehouse,
            condition=PhysicalReturnReceipt.Condition.SELLABLE,
            recorded_by=self.user,
        )
        damaged_receipt = PhysicalReturnReceipt.objects.create(
            sales_line=damaged_line,
            received_date=date(2026, 9, 21),
            quantity=1,
            warehouse=warehouse,
            condition=PhysicalReturnReceipt.Condition.DAMAGED,
            recorded_by=self.user,
        )
        InventoryMovement.objects.create(
            movement_key="FINANCE-RETURN-JOURNAL-MOVEMENT",
            movement_date=sellable_receipt.received_date,
            movement_type=InventoryMovement.MovementType.RETURN_IN,
            direction=InventoryMovement.Direction.IN,
            sku=sellable_sku,
            warehouse=warehouse,
            quantity=2,
            allocated_cost=Decimal("120000"),
            source_reference="FINANCE-RETURN-JOURNAL-001",
            return_receipt=sellable_receipt,
            posted_by=self.user,
        )
        self._allocate_sales_line(sellable_line, "110301")
        self._allocate_sales_line(damaged_line, "110101")
        params = {
            "start_date": date(2026, 9, 1),
            "end_date": date(2026, 9, 30),
            "actor": self.user,
        }

        preview = sales_return_journal_preview(
            start_date=params["start_date"],
            end_date=params["end_date"],
        )
        entry = create_sales_return_journal_draft(**params)

        self.assertEqual(preview["receipt_count"], 2)
        self.assertEqual(preview["receipt_quantity"], 3)
        self.assertEqual(preview["transaction_count"], 1)
        self.assertEqual(
            {line["account_code"] for line in preview["lines"]},
            {"440103", "110301", "110101", "110401", "5101"},
        )
        self.assertEqual(entry.status, JournalEntry.Status.DRAFT)
        self.assertEqual(entry.debit_total, Decimal("380000"))
        self.assertEqual(entry.credit_total, Decimal("380000"))
        lines = {line.account.code: line for line in entry.lines.select_related("account")}
        self.assertEqual(lines["440103"].debit, Decimal("260000"))
        self.assertEqual(lines["110301"].credit, Decimal("180000"))
        self.assertEqual(lines["110101"].credit, Decimal("80000"))
        self.assertEqual(lines["110401"].debit, Decimal("120000"))
        self.assertEqual(lines["5101"].credit, Decimal("120000"))
        self.assertEqual(SalesReturnJournalAllocation.objects.filter(entry=entry).count(), 2)
        self.assertTrue(
            SalesReturnJournalAllocation.objects.filter(
                return_receipt=damaged_receipt,
                reversed_cogs=0,
            ).exists()
        )
        with self.assertRaisesMessage(ValidationError, "belum dijurnal"):
            create_sales_return_journal_draft(**params)

        self.client.force_login(self.user)
        profit_loss = self.client.get(
            reverse("finance:profit_loss"),
            {"start": "2026-09-01", "end": "2026-09-30"},
        )
        revenue = {row["account"].code: row["amount"] for row in profit_loss.context["revenue"]}
        expense = {row["account"].code: row["amount"] for row in profit_loss.context["expense"]}
        self.assertEqual(revenue["440103"], Decimal("-260000"))
        self.assertEqual(expense["5101"], Decimal("-120000"))
        self.assertEqual(profit_loss.context["profit"], Decimal("120000"))

    def test_profit_loss_includes_draft_and_posted_sales_journal(self):
        from sales.models import SalesOrder, SalesOrderLine

        sku = self._sales_sku("FIN-PL-SKU")
        order = SalesOrder.objects.create(
            source=SalesOrder.Source.SHOPEE,
            source_label="Shopee",
            order_number="FIN-PL-ORDER",
            order_datetime=timezone.make_aware(datetime(2026, 9, 10, 10, 0)),
            order_date=date(2026, 9, 10),
            current_status="Retur",
            source_status="Retur",
            is_final=True,
            first_seen_batch_id=uuid.uuid4(),
            latest_batch_id=uuid.uuid4(),
        )
        SalesOrderLine.objects.create(
            order=order,
            sku=sku,
            sku_code_snapshot="FIN-PL-SKU",
            category_snapshot="Finance T-Shirt",
            product_name_snapshot="Finance P&L Product",
            current_status="Selesai",
            is_final=True,
            quantity=2,
            net_unit_price=Decimal("90000"),
            retail_price_snapshot=Decimal("100000"),
            total_gross_sales=Decimal("200000"),
            total_net_sales=Decimal("180000"),
            total_cogs=Decimal("120000"),
        )
        self.client.force_login(self.user)

        before = self.client.get(
            reverse("finance:profit_loss"),
            {"start": "2026-09-01", "end": "2026-09-30"},
        )
        self.assertEqual(before.context["gross_sales_total"], Decimal("0"))
        self.assertEqual(before.context["profit"], Decimal("0"))
        self.assertContains(before, "Belum ada jurnal Draft atau Posted")

        entry = create_sales_journal_draft(
            start_date=date(2026, 9, 1),
            end_date=date(2026, 9, 30),
            receipt_account_id=Account.objects.get(code="110301").id,
            actor=self.user,
        )
        draft = self.client.get(
            reverse("finance:profit_loss"),
            {"start": "2026-09-01", "end": "2026-09-30"},
        )
        self.assertEqual(draft.context["gross_sales_total"], Decimal("200000"))
        self.assertEqual(draft.context["revenue_total"], Decimal("180000"))
        self.assertEqual(draft.context["expense_total"], Decimal("120000"))
        post_journal(entry.id, self.user)
        response = self.client.get(
            reverse("finance:profit_loss"),
            {"start": "2026-09-01", "end": "2026-09-30"},
        )

        revenue = {row["account"].code: row["amount"] for row in response.context["revenue"]}
        expense = {row["account"].code: row["amount"] for row in response.context["expense"]}
        self.assertEqual(revenue["410002"], Decimal("200000"))
        self.assertEqual(revenue["440101"], Decimal("-20000"))
        self.assertEqual(expense["5101"], Decimal("120000"))
        self.assertEqual(response.context["gross_sales_total"], Decimal("200000"))
        self.assertEqual(response.context["revenue_total"], Decimal("180000"))
        self.assertEqual(response.context["expense_total"], Decimal("120000"))
        self.assertEqual(response.context["profit"], Decimal("60000"))
        self.assertContains(response, "Beban Pokok Penjualan")
        self.assertContains(response, "Subtotal Gross Sales")
        self.assertContains(response, "Subtotal Diskon Penjualan")
        self.assertContains(response, "Subtotal Beban Pokok Penjualan")
        self.assertContains(response, "Total Pendapatan Bersih")
        self.assertContains(response, "<th>Nilai</th>", html=True)
        self.assertNotContains(response, "Kelompok")

        category_view = self.client.get(
            reverse("finance:profit_loss"),
            {"start": "2026-09-01", "end": "2026-09-30", "sales_view": "category"},
        )
        category_row = category_view.context["sales_dimension_rows"][0]
        self.assertEqual(category_view.context["sales_view"], "category")
        self.assertEqual(category_row["label"], "Finance T-Shirt")
        self.assertEqual(category_row["gross"], Decimal("200000"))
        self.assertEqual(category_row["discount"], Decimal("20000"))
        self.assertEqual(category_row["net"], Decimal("180000"))
        self.assertEqual(category_row["cogs"], Decimal("120000"))
        self.assertContains(category_view, "Profit &amp; Loss per Kategori")
        self.assertContains(category_view, "Total Gross Sales")
        self.assertContains(category_view, "Diskon Penjualan")
        self.assertContains(category_view, "Net Sales")
        self.assertContains(category_view, "Total COGS")
        self.assertContains(category_view, "Laba Kotor")
        self.assertContains(category_view, "GPM")
        self.assertNotContains(category_view, "Total Diskon Penjualan")
        self.assertNotContains(category_view, "Total Net Sales")
        self.assertNotContains(category_view, "Total Laba Kotor")
        self.assertNotContains(category_view, "Total GPM")

        source_view = self.client.get(
            reverse("finance:profit_loss"),
            {"start": "2026-09-01", "end": "2026-09-30", "sales_view": "source"},
        )
        source_row = source_view.context["sales_dimension_rows"][0]
        self.assertEqual(source_row["label"], "Shopee")
        self.assertEqual(source_view.context["sales_dimension_total"]["gross_profit"], Decimal("60000"))
        self.assertContains(source_view, "Profit &amp; Loss per Source")

        multi_period = self.client.get(
            reverse("finance:profit_loss"),
            {"mode": "multi_period", "start_month": "2026-08", "end_month": "2026-09"},
        )
        self.assertEqual(
            [period["label"] for period in multi_period.context["comparison"]["periods"]],
            ["Aug 2026", "Sep 2026"],
        )
        self.assertEqual(multi_period.context["comparison"]["profits"], [Decimal("0"), Decimal("60000")])
        self.assertContains(multi_period, "Profit &amp; Loss Multi Period")

        multi_year = self.client.get(
            reverse("finance:profit_loss"),
            {"mode": "multi_year", "year": "2026"},
        )
        self.assertEqual(
            [period["label"] for period in multi_year.context["comparison"]["periods"]],
            ["2024", "2025", "2026"],
        )
        self.assertEqual(multi_year.context["comparison"]["profits"], [Decimal("0"), Decimal("0"), Decimal("60000")])
        self.assertContains(multi_year, "Profit &amp; Loss Multi Year")

    def test_finance_workspace_respects_exact_tab_permission(self):
        user = get_user_model().objects.create_user(
            username="cashier",
            password="test",
            module_access={"finance": "view"},
            tab_access={"finance": ["other_payment"]},
        )
        self.client.force_login(user)

        self.assertEqual(self.client.get(reverse("finance:feature", args=["other-payment"])).status_code, 200)
        self.assertEqual(self.client.get(reverse("finance:feature", args=["other-deposit"])).status_code, 403)

    def test_sales_settings_filters_and_saves_product_revenue_account(self):
        regular = ProductStatus.objects.create(code="FIN-REG", name="Regular Finance")
        seasonal = ProductStatus.objects.create(code="FIN-SEA", name="Seasonal Finance")
        shirts = Category.objects.create(code="FIN-SHIRT", name="Finance Shirts")
        pants = Category.objects.create(code="FIN-PANTS", name="Finance Pants")
        oxford = Subcategory.objects.create(category=shirts, code="FIN-OXF", name="Finance Oxford")
        product = Product.objects.create(
            code="FIN-PRODUCT-1",
            parent_sku="FIN-PARENT-1",
            name="Finance Product One",
            status=regular,
            category=shirts,
            subcategory=oxford,
        )
        Product.objects.create(
            code="FIN-PRODUCT-2",
            parent_sku="FIN-PARENT-2",
            name="Finance Product Two",
            status=seasonal,
            category=pants,
        )
        revenue = Account.objects.get(code="410001")
        self.client.force_login(self.user)

        page = self.client.get(reverse("finance:sales_settings"), {"product_status": regular.id})

        self.assertContains(page, "Finance Product One")
        self.assertNotContains(page, "Finance Product Two")
        self.assertContains(page, "Finance Shirts")
        self.assertNotContains(page, "Finance Pants")

        response = self.client.post(
            f"{reverse('finance:sales_settings')}?product_status={regular.id}",
            {"product_id": product.id, "sales_account_id": revenue.id},
        )

        self.assertRedirects(response, f"{reverse('finance:sales_settings')}?product_status={regular.id}")
        self.assertEqual(ProductSalesAccount.objects.get(product=product).sales_account, revenue)
        self.assertTrue(
            AuditEvent.objects.filter(
                action="finance_sales_account_mapping_updated",
                entity_id=product.id,
            ).exists()
        )

    def test_sales_settings_viewer_cannot_change_mapping(self):
        viewer = get_user_model().objects.create_user(
            username="finance-settings-viewer",
            password="test",
            module_access={"finance": "view"},
            tab_access={"finance": ["sales_settings"]},
        )
        self.client.force_login(viewer)

        self.assertEqual(self.client.get(reverse("finance:sales_settings")).status_code, 200)
        self.assertEqual(self.client.post(reverse("finance:sales_settings"), {}).status_code, 403)

    def test_sales_settings_bulk_updates_selected_products(self):
        regular = ProductStatus.objects.create(code="FIN-BULK", name="Regular Finance Bulk")
        category = Category.objects.create(code="FIN-BULK-CAT", name="Finance Bulk Category")
        products = [
            Product.objects.create(
                code=f"FIN-BULK-{index}",
                parent_sku=f"FIN-BULK-PARENT-{index}",
                name=f"Finance Bulk Product {index}",
                status=regular,
                category=category,
            )
            for index in range(1, 3)
        ]
        revenue = Account.objects.get(code="410001")
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("finance:sales_settings"),
            {
                "action": "bulk_update",
                "product_ids": [product.id for product in products],
                "sales_account_id": revenue.id,
            },
        )

        self.assertRedirects(response, reverse("finance:sales_settings"))
        self.assertEqual(
            ProductSalesAccount.objects.filter(product__in=products, sales_account=revenue).count(),
            2,
        )
        self.assertEqual(
            AuditEvent.objects.filter(
                action="finance_sales_account_mapping_updated",
                entity_id__in=[product.id for product in products],
            ).count(),
            2,
        )
        page = self.client.get(reverse("finance:sales_settings"), {"q": "Finance Bulk Product"})
        self.assertContains(page, "Ubah Massal")
        self.assertContains(page, "data-sales-select-all")

    def test_cash_workspace_prefills_matching_journal_workflow(self):
        self.client.force_login(self.user)

        response = self.client.get(reverse("finance:journal_create"), {"workflow": "bank-transfer"})

        self.assertContains(response, 'name="workflow" value="bank-transfer"')
        self.assertEqual(response.context["form"].initial["description"], "Bank Transfer")

    def test_created_journal_is_marked_as_uat(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("finance:journal_create"),
            {
                "entry_date": "2026-10-10",
                "description": "UAT payment",
                "reference": "UAT-001",
                "lines-TOTAL_FORMS": "2",
                "lines-INITIAL_FORMS": "0",
                "lines-MIN_NUM_FORMS": "0",
                "lines-MAX_NUM_FORMS": "1000",
                "lines-0-account": str(self.cash.id),
                "lines-0-description": "",
                "lines-0-debit": "100",
                "lines-0-credit": "0",
                "lines-1-account": str(self.capital.id),
                "lines-1-description": "",
                "lines-1-debit": "0",
                "lines-1-credit": "100",
            },
        )
        self.assertEqual(response.status_code, 302)
        created = JournalEntry.objects.filter(source=JournalEntry.Source.MANUAL).exclude(pk=self.entry.pk).get()
        self.assertEqual(created.source_metadata["environment"], "UAT")

    def test_account_page_shows_opening_balance_for_leaf_and_parent(self):
        self.client.force_login(self.user)

        response = self.client.get(reverse("finance:accounts"), {"q": "1101"})

        rows = {account.code: account for account in response.context["accounts"]}
        self.assertEqual(rows["110101"].opening_debit, Decimal("15181"))
        self.assertGreater(rows["1101"].opening_debit, rows["110101"].opening_debit)
        self.assertContains(response, "Saldo Awal")

    def test_superadmin_can_edit_coa_with_audit_history(self):
        self.client.force_login(self.user)
        previous_name = self.cash.name

        response = self.client.post(
            reverse("finance:accounts"),
            {
                "account_id": str(self.cash.id),
                "code": self.cash.code,
                "name": "Petty Cash UAT",
                "account_type": self.cash.account_type,
                "parent": str(self.cash.parent_id),
                "currency": "IDR",
                "is_postable": "on",
                "is_active": "on",
                "opening_balance": "15181",
                "opening_side": "DEBIT",
            },
        )

        self.assertRedirects(response, reverse("finance:accounts"))
        self.cash.refresh_from_db()
        self.assertEqual(self.cash.name, "Petty Cash UAT")
        audit = AuditEvent.objects.get(action="finance_account_updated", entity_id=self.cash.id)
        self.assertEqual(audit.before_values["name"], previous_name)
        self.assertEqual(audit.after_values["name"], "Petty Cash UAT")

    def test_new_coa_can_set_balanced_opening_balance(self):
        self.client.force_login(self.user)
        form_page = self.client.get(reverse("finance:accounts"), {"add": "1"})
        self.assertContains(form_page, 'name="opening_balance"')
        self.assertContains(form_page, 'name="opening_side"')
        self.assertContains(form_page, 'name="is_subaccount"')
        self.assertNotContains(form_page, 'name="is_postable"')

        response = self.client.post(
            reverse("finance:accounts"),
            {
                "code": "990002",
                "name": "UAT Opening Account",
                "account_type": "BANK",
                "is_subaccount": "on",
                "parent": str(self.cash.parent_id),
                "currency": "IDR",
                "is_active": "on",
                "opening_balance": "250000",
                "opening_side": "DEBIT",
            },
        )

        self.assertRedirects(response, reverse("finance:accounts"))
        account = Account.objects.get(code="990002")
        opening = JournalEntry.objects.get(number="OPENING-20260831")
        line = opening.lines.get(account=account)
        offset = opening.lines.get(account__code="300001")
        self.assertEqual(line.debit, Decimal("250000"))
        self.assertEqual(offset.credit, Decimal("250000"))
        self.assertEqual(opening.debit_total, opening.credit_total)
        self.assertEqual(opening.status, JournalEntry.Status.DRAFT)
        self.assertTrue(account.is_postable)
        self.assertEqual(account.parent_id, self.cash.parent_id)
        self.assertTrue(
            AuditEvent.objects.filter(action="finance_opening_balance_updated", entity_id=account.id).exists()
        )

    def test_new_coa_without_subaccount_is_created_as_parent(self):
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("finance:accounts"),
            {
                "code": "990003",
                "name": "UAT Parent Account",
                "account_type": "OASS",
                "parent": str(self.cash.parent_id),
                "currency": "IDR",
                "is_active": "on",
                "opening_balance": "0",
                "opening_side": "DEBIT",
            },
        )

        self.assertRedirects(response, reverse("finance:accounts"))
        account = Account.objects.get(code="990003")
        self.assertFalse(account.is_postable)
        self.assertIsNone(account.parent_id)

    def test_new_subaccount_requires_parent(self):
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("finance:accounts"),
            {
                "code": "990004",
                "name": "UAT Child Without Parent",
                "account_type": "OASS",
                "is_subaccount": "on",
                "currency": "IDR",
                "is_active": "on",
                "opening_balance": "0",
                "opening_side": "DEBIT",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Pilih akun induk untuk membuat Sub-account.")
        self.assertFalse(Account.objects.filter(code="990004").exists())

    def test_edit_coa_prefills_its_direct_opening_balance(self):
        self.client.force_login(self.user)

        response = self.client.get(reverse("finance:accounts"), {"edit": self.cash.id})

        self.assertEqual(response.context["account_form"]["opening_balance"].value(), Decimal("15181.000000"))
        self.assertEqual(response.context["account_form"]["opening_side"].value(), "DEBIT")

    def test_used_account_cannot_be_changed_into_parent(self):
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("finance:accounts"),
            {
                "account_id": str(self.cash.id),
                "code": self.cash.code,
                "name": self.cash.name,
                "account_type": self.cash.account_type,
                "parent": str(self.cash.parent_id),
                "currency": "IDR",
                "is_active": "on",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "sudah dipakai jurnal")
        self.cash.refresh_from_db()
        self.assertTrue(self.cash.is_postable)

    def test_parent_account_cannot_be_changed_into_transaction_account(self):
        self.client.force_login(self.user)
        parent = self.cash.parent

        response = self.client.post(
            reverse("finance:accounts"),
            {
                "account_id": str(parent.id),
                "code": parent.code,
                "name": parent.name,
                "account_type": parent.account_type,
                "currency": "IDR",
                "is_postable": "on",
                "is_active": "on",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "memiliki akun turunan")
        parent.refresh_from_db()
        self.assertFalse(parent.is_postable)


class FinanceCutoverImportTests(TestCase):
    def setUp(self):
        JournalLine.objects.all().delete()
        JournalEntry.objects.all().delete()
        Account.objects.update(parent=None)
        Account.objects.all().delete()

    def _write_sources(self, directory):
        coa_path = Path(directory) / "coa.xlsx"
        coa = Workbook()
        sheet = coa.active
        sheet.append(["No.", "Tipe Akun", "Kode Perkiraan", "Nama", "Akun Induk", "Mata Uang"])
        sheet.append([1, "OASS", "1000", "Assets", "", "IDR"])
        sheet.append([2, "BANK", "1001", "Cash", "1000", "IDR"])
        sheet.append([3, "EQTY", "3001", "Capital", "", "IDR"])
        coa.save(coa_path)

        tb_path = Path(directory) / "tb.xlsx"
        tb = Workbook()
        sheet = tb.active
        sheet.cell(5, 2, "Kode")
        sheet.cell(5, 4, "Nama")
        sheet.cell(5, 6, "Saldo Awal (Debit)")
        sheet.cell(5, 8, "Saldo Awal (Kredit)")
        sheet.cell(5, 10, "Perubahan Debit")
        sheet.cell(5, 12, "Perubahan Kredit")
        sheet.cell(5, 14, "Saldo Akhir (Debit)")
        sheet.cell(5, 16, "Saldo Akhir (Kredit)")
        for row_number, code, name, debit, credit in (
            (6, "1000", "Assets", 100, 0),
            (7, "1001", "Cash", 100, 0),
            (8, "3001", "Capital", 0, 100),
        ):
            sheet.cell(row_number, 2, code)
            sheet.cell(row_number, 4, name)
            sheet.cell(row_number, 10, debit)
            sheet.cell(row_number, 12, credit)
            sheet.cell(row_number, 14, debit)
            sheet.cell(row_number, 16, credit)
        tb.save(tb_path)
        return coa_path, tb_path

    def test_cutover_import_stages_balanced_opening_without_posting(self):
        with TemporaryDirectory() as directory:
            coa_path, tb_path = self._write_sources(directory)
            entry = stage_finance_cutover(
                coa_path=coa_path,
                trial_balance_path=tb_path,
                cutoff_date=date(2026, 8, 31),
            )
        self.assertEqual(Account.objects.count(), 3)
        self.assertFalse(Account.objects.get(code="1000").is_postable)
        self.assertEqual(entry.status, JournalEntry.Status.DRAFT)
        self.assertEqual(entry.entry_date, date(2026, 9, 1))
        self.assertEqual(entry.lines.count(), 2)
        self.assertEqual(entry.debit_total, entry.credit_total)
        self.assertEqual(entry.source_metadata["cutoff_date"], "2026-08-31")
